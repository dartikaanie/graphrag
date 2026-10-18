"""Amendment 4 tooling: binary label, v2.1 runner, calibration guard, blind annotation export/import,
agreement. Offline -- no LLM call, no parquet, no real key directory."""

import json
from pathlib import Path

import pytest

import _judge_common
import annotation_tool as at
import judge_binary as jb
import llm_judge_hallucination_v2 as runner_v2
import llm_judge_hallucination_v2_1 as runner_v21


def _claim(verdict, severity=None):
    return {"claim": "c", "verdict": verdict, "severity": severity, "reference_conflict": False, "evidence": "e"}


def _rec(attempted=True, claims=(), **kw):
    return {"answer_attempted": attempted, "claims": list(claims), "parse_error": False, "call_failed": False, **kw}


# ---------------------------------------------------------------- binary label

def test_binary_derivation():
    assert jb.derive_hallucinated(_rec(claims=[_claim("SUPPORTED"), _claim("UNVERIFIABLE")])) is False
    assert jb.derive_hallucinated(_rec(claims=[])) is False
    assert jb.derive_hallucinated(_rec(claims=[_claim("CONTRADICTED", "MINOR")])) is True
    assert jb.derive_hallucinated(_rec(claims=[_claim("SUPPORTED"), _claim("FABRICATED", "CORE")])) is True
    assert jb.derive_hallucinated(_rec(claims=[_claim("FABRICATED", None)])) is True      # any severity, even none
    # ABSTAIN / parse error / failed call -> null
    assert jb.derive_hallucinated(_rec(attempted=False, claims=[_claim("FABRICATED", "CORE")])) is None
    assert jb.derive_hallucinated({"parse_error": True, "claims": []}) is None
    assert jb.derive_hallucinated({"call_failed": True, "parse_error": True, "answer_attempted": None}) is None


def test_category_and_severity():
    assert jb.category_of(_rec(attempted=False)) == "ABSTAIN"
    assert jb.category_of(_rec(claims=[_claim("UNVERIFIABLE")])) == "NOT_HALLUCINATED"
    assert jb.category_of(_rec(claims=[_claim("CONTRADICTED", "CORE")])) == "HALLUCINATED"
    assert jb.category_of({"parse_error": True}) is None
    assert jb.severity_of(_rec(claims=[_claim("CONTRADICTED", "MINOR"), _claim("FABRICATED", "CORE")])) == "PENUH"
    assert jb.severity_of(_rec(claims=[_claim("CONTRADICTED", "MINOR")])) == "SEBAGIAN"
    assert jb.severity_of(_rec(claims=[_claim("CONTRADICTED", "MINOR"), _claim("FABRICATED", None)])) == "UNSPECIFIED"
    assert jb.severity_of(_rec(claims=[_claim("SUPPORTED")])) is None


def test_binary_matches_v2_three_class_label_on_real_pilot_records():
    """For judge-v2, hallucinated == (label is HALUSINASI_*) and ABSTAIN -> null, on every real pilot record."""
    path = Path(__file__).resolve().parents[1] / "results" / "jv2_dashboard_primary_judge-v2_reparsed.jsonl"
    records = [json.loads(line) for line in open(path)]
    assert len(records) == 90
    for r in records:
        expected = None if r["label"] == "ABSTAIN" else r["label"] in ("HALUSINASI_SEBAGIAN", "HALUSINASI_PENUH")
        assert jb.hallucinated_of(r) is expected
        assert "hallucinated" not in r                    # v2 files are not rewritten


def test_binary_summary_rates():
    recs = ([_rec(claims=[_claim("CONTRADICTED", "CORE"), _claim("UNVERIFIABLE")])] * 2
            + [_rec(claims=[_claim("FABRICATED", "MINOR")])] + [_rec(claims=[_claim("FABRICATED", None)])]
            + [_rec(claims=[_claim("SUPPORTED"), _claim("UNVERIFIABLE")])] * 4 + [_rec(attempted=False)] * 2
            + [{"parse_error": True}])
    s = jb.binary_summary(recs)
    assert (s["n_records"], s["n_unusable"], s["n"]) == (11, 1, 10)
    assert s["hallucination_rate_excl_abstain"] == 0.5          # 4 of 8 attempted
    assert s["hallucination_rate_abstain_counted"] == 0.6       # (4 + 2) of 10
    assert s["abstain_rate"] == 0.2
    assert s["severity_distribution"] == {"PENUH": 2, "SEBAGIAN": 1, "UNSPECIFIED": 1}
    assert s["n_claims"] == 14 and s["unverifiable_rate"] == round(6 / 14, 4)


def test_kappa_on_a_known_example():
    # Classic 2x2: 20 both-yes, 5 yes/no, 10 no/yes, 15 both-no -> po = 0.7, pe = 0.5, kappa = 0.4
    a = ["Y"] * 25 + ["N"] * 25
    b = ["Y"] * 20 + ["N"] * 5 + ["Y"] * 10 + ["N"] * 15
    assert jb.cohen_kappa(a, b, ("Y", "N")) == pytest.approx(0.4)
    assert jb.cohen_kappa(["Y", "N"], ["Y", "N"], ("Y", "N")) == pytest.approx(1.0)
    assert jb.cohen_kappa(["Y", "Y"], ["Y", "Y"], ("Y", "N")) is None      # undefined, not 1.0
    assert jb.cohen_kappa([], [], ("Y", "N")) is None


def test_agreement_report_binary_kappa_only_on_both_non_abstain():
    A, N, H = jb.CATEGORIES
    human = [A, A, N, N, N, H, H, H, H, N]
    judge = [A, N, N, N, H, H, H, H, N, A]
    rep = jb.agreement_report(human, judge)
    assert rep["n"] == 10 and rep["n_both_non_abstain"] == 7      # drops the 3 items either one calls ABSTAIN
    assert rep["confusion_3cat"][A][N] == 1 and rep["confusion_3cat"][N][A] == 1
    assert rep["confusion_binary_non_abstain"] == {N: {N: 2, H: 1}, H: {N: 1, H: 3}}
    assert rep["kappa_binary_non_abstain"] == pytest.approx((5 / 7 - (3 / 7 * 3 / 7 + 4 / 7 * 4 / 7)) /
                                                             (1 - (3 / 7 * 3 / 7 + 4 / 7 * 4 / 7)), abs=1e-4)


# ---------------------------------------------------------------- runners

class _FakeConfig:
    model = "fake-model"
    base_url = None


def _item():
    return {"run_id": "A-1", "condition": "A", "run_label": "A [dev]", "fusion_mode": None, "retrieval_version": None,
            "source_file": "/fake.jsonl", "question_id": 1, "title": "t", "tags": "", "candidate": "c",
            "ground_truth_answer_html": "ref"}


def _payload(claims, attempted=True, label="FAKTUAL", finish="stop", request_params=None):
    return {"raw": json.dumps({"answer_attempted": attempted, "completeness": "FULL", "claims": claims,
                               "reasoning": "r", "label": label}),
            "prompt_tokens": 1, "completion_tokens": 1, "latency_s": 0.1, "response_id": "x",
            "response_model": "fake-model", "response_created": 1, "system_fingerprint": None,
            "finish_reason": finish, "usage_full": {}, "request_params": request_params or {}}


def test_v2_1_runner_records_version_settings_and_binary_label(monkeypatch):
    sent = {}

    def fake_call(client, config, messages, temperature, log, max_tokens):
        sent.update(temperature=temperature, max_tokens=max_tokens, system=messages[0]["content"])
        return _payload([_claim("FABRICATED", "CORE"), _claim("SUPPORTED", "MINOR")], label="HALUSINASI_PENUH",
                        request_params={"model": "fake-model", "temperature": temperature, "max_tokens": max_tokens}), []

    monkeypatch.setattr(runner_v21, "call_judge_llm_with_retry", fake_call)
    rec = runner_v21.judge_one(_item(), "primary", None, _FakeConfig(), "body", "ref")
    assert sent["temperature"] == 0.0 and sent["max_tokens"] == 1500
    assert "judge-v2.1" not in sent["system"] and "Keep the qualifiers" in sent["system"]
    assert (rec["prompt_version"], rec["judge_version"], rec["parse_version"]) == ("judge-v2.1", "judge-v2.1", "pv2.1-1")
    assert rec["temperature"] == 0.0 and rec["max_tokens"] == 1500
    assert rec["hallucinated"] is True and rec["severity_label"] == "PENUH" and rec["label"] == "HALUSINASI_PENUH"
    assert rec["normalized_fields"] == ["severity_on_non_error"] and rec["truncated"] is False
    assert _judge_common.judge_resume_key(rec)[5] == "judge-v2.1"

    monkeypatch.setattr(runner_v21, "call_judge_llm_with_retry",
                        lambda *a: (_payload([], attempted=False, label="ABSTAIN"), []))
    assert runner_v21.judge_one(_item(), "primary", None, _FakeConfig(), "b", "r")["hallucinated"] is None
    monkeypatch.setattr(runner_v21, "call_judge_llm_with_retry", lambda *a: (None, [{"error": "boom"}]))
    failed = runner_v21.judge_one(_item(), "primary", None, _FakeConfig(), "b", "r")
    assert failed["call_failed"] is True and failed["hallucinated"] is None and failed["max_tokens"] == 1500


def test_v2_runner_output_unchanged(monkeypatch):
    """judge-v2 stays as it was: no cap sent, no v2.1 fields, same prompt and parser."""
    sent = {}

    def fake_call(client, config, messages, temperature, log, max_tokens):
        sent.update(max_tokens=max_tokens, system=messages[0]["content"])
        return _payload([_claim("CONTRADICTED", "MINOR")], label="HALUSINASI_SEBAGIAN",
                        request_params={"model": "fake-model", "temperature": 0.0}), []

    monkeypatch.setattr(runner_v2, "call_judge_llm_with_retry", fake_call)
    rec = runner_v2.judge_one(_item(), "primary", None, _FakeConfig(), "body", "ref")
    import judge_prompt_v2

    assert sent["max_tokens"] is None and sent["system"] == judge_prompt_v2.SYSTEM_PROMPT
    assert rec["prompt_version"] == "judge-v2" and rec["parse_version"] == "pv1" and rec["max_tokens"] is None
    assert rec["label"] == "HALUSINASI_SEBAGIAN"
    for new_field in ("hallucinated", "severity_label", "judge_version", "temperature", "n_claims"):
        assert new_field not in rec
    assert set(rec) == {
        "run_id", "config_hash", "condition", "run_label", "fusion_mode", "grounding", "retrieval_version",
        "source_file", "question_id", "judge_id", "judge_model", "judge_base_url", "prompt_version",
        "blinding_version", "timestamp_utc", "messages_sha256", "attempt_count", "attempts", "label",
        "derived_label", "judge_reported_label", "consistent", "answer_attempted", "completeness",
        "attempted_claims_conflict", "parse_error", "claims", "reasoning", "raw_response", "prompt_tokens",
        "completion_tokens", "latency_s", "call_failed", "error", "response_id", "response_model",
        "response_created", "system_fingerprint", "finish_reason", "usage_full", "request_params",
        "parse_version", "normalized_fields", "max_tokens", "truncated", "served_model_mismatch"}


# ---------------------------------------------------------------- calibration guard

def _registry(tmp_path, run_file):
    reg = tmp_path / "calibration_runs.json"
    rel = Path(run_file).resolve().relative_to(_judge_common._EVAL_DIR.parents[1])
    reg.write_text(json.dumps({"n_annotation_items": 2, "runs": [{"output_path": rel.as_posix()}]}))
    return reg


def test_calibration_guard(tmp_path):
    cal = _judge_common._EVAL_DIR / "results" / "_guard_test_calibration_run.jsonl"   # path only; never created
    other = tmp_path / "other.jsonl"
    reg, ref = _registry(tmp_path, cal), tmp_path / "reference_adjudicated.json"
    guard = lambda files, override=False, logger=print: _judge_common.assert_calibration_judging_allowed(  # noqa: E731
        files, override, logger, registry_path=reg, reference_path=ref)

    assert guard([other]) == {"calibration_run_files": [], "reference_ok": None, "reference_status": None,
                              "override_used": False}
    with pytest.raises(_judge_common.CalibrationGuardError, match="adjudicated reference not found"):
        guard([other, cal])

    logged = []
    info = guard([cal], override=True, logger=logged.append)
    assert info["override_used"] is True and len(logged) == 1 and "CALIBRATION GUARD OVERRIDDEN" in logged[0]

    ref.write_text(json.dumps({"status": "draft", "items": [1, 2]}))
    with pytest.raises(_judge_common.CalibrationGuardError, match="not 'adjudicated'"):
        guard([cal])
    ref.write_text(json.dumps({"status": "adjudicated", "items": [1]}))
    with pytest.raises(_judge_common.CalibrationGuardError, match="expected 2"):
        guard([cal])
    ref.write_text(json.dumps({"status": "adjudicated", "items": [1, 2]}))
    assert guard([cal])["reference_ok"] is True and guard([cal])["override_used"] is False


@pytest.mark.parametrize("runner", [runner_v2, runner_v21])
def test_both_judge_runners_call_the_guard_before_reading_any_run_file(monkeypatch, runner):
    def refuse(run_files, allow_override, log):
        raise _judge_common.CalibrationGuardError(f"refused override={allow_override}")

    monkeypatch.setattr(_judge_common, "assert_calibration_judging_allowed", refuse)
    with pytest.raises(_judge_common.CalibrationGuardError, match="override=False"):
        runner.load_items(["/does/not/exist.jsonl"], None)
    with pytest.raises(_judge_common.CalibrationGuardError, match="override=True"):
        runner.load_items(["/does/not/exist.jsonl"], None, allow_unannotated_calibration=True)


# ---------------------------------------------------------------- sampling + blind export

def _items():
    items = []
    for label, cond in (("A [dev]", "A"), ("B-plain [dev]", "B"), ("B-grounded [dev]", "B")):
        for q in range(1000, 1050):
            items.append({"run_id": f"{cond}-{label[:6]}", "run_label": label, "condition": cond,
                          "source_file": f"/x/{label}.jsonl", "config_hash": "h", "question_id": q,
                          "title": f"Title {q}", "ground_truth_answer_html": f"<p>Reference {q}</p>",
                          "candidate": f"Answer for {q} per the docs [SO-{q}] and also [SO-9]. Sources: [SO-{q}]"})
    return items


def test_calibration_sampling_is_stratified_and_deterministic():
    s1 = at.sample_items(_items(), "calibration", 42)
    s2 = at.sample_items(list(reversed(_items())), "calibration", 42)        # input order must not matter
    assert [(i["item_id"], i["run_label"], i["question_id"], i["in_annotator2_subset"]) for i in s1] == \
           [(i["item_id"], i["run_label"], i["question_id"], i["in_annotator2_subset"]) for i in s2]
    assert len(s1) == 60 and [i["item_id"] for i in s1] == [f"CAL-{n:03d}" for n in range(1, 61)]
    per = {}
    for i in s1:
        per.setdefault(i["run_label"], []).append(i)
    assert {k: len(v) for k, v in per.items()} == {"A [dev]": 20, "B-plain [dev]": 20, "B-grounded [dev]": 20}
    assert {k: sum(x["in_annotator2_subset"] for x in v) for k, v in per.items()} == \
           {"A [dev]": 10, "B-plain [dev]": 10, "B-grounded [dev]": 10}
    assert [i["question_id"] for i in s1] != [i["question_id"] for i in at.sample_items(_items(), "calibration", 7)]
    # item ids are assigned after the shuffle: the first 20 ids are not one condition
    assert len({i["run_label"] for i in s1[:20]}) > 1
    a2 = at.annotator2_order(s1, "calibration", 42)
    assert len(a2) == 30 and [i["item_id"] for i in a2] != sorted(i["item_id"] for i in a2)


def test_confirmation_sampling_is_7_per_official_run():
    """Amendment 11: 7 items per stratum (the nine official runs); allocate_evenly stays for other uses."""
    cells = ["A", "B-plain", "B-grounded", "C-primary-plain", "C-primary-grounded", "C-ablation-plain",
             "C-ablation-grounded", "D-plain", "D-grounded"]
    items = [dict(i, run_label=c) for c in cells for i in _items()[:50]]
    s = at.sample_items(items, "confirmation", 42)
    assert len(s) == 63 and s[0]["item_id"] == "CONF-001" and not any(i["in_annotator2_subset"] for i in s)
    assert {sum(1 for i in s if i["run_label"] == c) for c in cells} == {7}
    assert at.allocate_evenly(["b", "a", "c", "d"], 30) == {"a": 8, "b": 8, "c": 7, "d": 7}
    with pytest.raises(at.AnnotationError):
        at.sample_items(items[:5], "confirmation", 42)


def _export(tmp_path, monkeypatch):
    monkeypatch.setattr(at, "ANNOTATION_ROOT", tmp_path / "repo" / "annotation")
    items = _items()
    bodies = {q: f"<p>Body of {q}</p>" for q in range(1000, 1050)}
    export = at.build_export(items, bodies, "calibration", 42)
    pdir, key_dir = at.purpose_dir("calibration"), tmp_path / "outside_repo_keys"
    key_dir.mkdir()
    key = {"purpose": "calibration", "seed": 42, "items": export["key_items"]}
    key_path = key_dir / "calibration_key_seed42.json"
    key_path.write_text(json.dumps(key))
    at.write_annotation_xlsx(pdir / "annotator1.xlsx", export["annotator1"])
    at.write_annotation_xlsx(pdir / "annotator2.xlsx", export["annotator2"])
    (pdir / "export_manifest.json").write_text(json.dumps({
        "key_file_name": key_path.name, "key_sha256": at.file_sha256(key_path), "n_items_annotator2": 30}))
    return export, key, key_dir, pdir


def test_blind_export_has_no_condition_or_judge_fields_and_no_citation_markers(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    export, key, key_dir, pdir = _export(tmp_path, monkeypatch)
    assert export["n_markers_stripped"] == 180 and all(k["citation_markers_stripped"] == 3 for k in key["items"])
    for name, n in (("annotator1.xlsx", 60), ("annotator2.xlsx", 30)):
        wb = load_workbook(pdir / name)
        ws = wb["anotasi"]
        assert [c.value for c in ws[1]] == at.COLUMNS and ws.max_row == n + 1 and ws.freeze_panes == "B2"
        assert ws["B2"].alignment.wrap_text is True
        assert sorted(dv.formula1 for dv in ws.data_validations.dataValidation) == sorted([
            '"Y,N"', '"Y,N"', '"PENUH,SEBAGIAN"', '"all_supported,partly_supported,not_supported"',
            '"same,different_valid,different_unassessable"', '"familiar,not_familiar"', '"Y,N"'])
        blob = "\n".join(str(c) for row in ws.iter_rows(values_only=True) for c in row if c is not None)
        blob += "\n".join(str(c) for row in wb["petunjuk"].iter_rows(values_only=True) for c in row if c)
        assert "[SO-" not in blob
        for forbidden in ("B-plain", "B-grounded", "[dev]", "run_label", "condition", "run_id", "judge",
                          "HALUSINASI", "SUPPORTED", "FABRICATED", "config_hash"):
            assert forbidden not in blob, forbidden
        assert "Answer for" in blob and "Body of" in blob and "Reference" in blob
    row = export["annotator1"][0]
    assert set(row) == {"item_id", "question", "accepted_answer", "candidate_answer", "context"}
    assert row["candidate_answer"].startswith("Answer for") and "per the docs" in row["candidate_answer"]


def test_key_is_verified_against_the_manifest_hash(tmp_path, monkeypatch):
    export, key, key_dir, pdir = _export(tmp_path, monkeypatch)
    assert at.load_verified_key("calibration", key_dir)["items"] == key["items"]
    (key_dir / "calibration_key_seed42.json").write_text(json.dumps({**key, "seed": 43}))
    with pytest.raises(at.AnnotationError, match="not the key of this export"):
        at.load_verified_key("calibration", key_dir)
    with pytest.raises(at.AnnotationError, match="key file not found"):
        at.load_verified_key("calibration", tmp_path / "nowhere")


def _fill(path, labels, edit=None):
    """labels: {item_id: (abstain, hallucinated, severity)}; rows not listed get ("N", "N", None). The other
    Amendment 11 columns are filled validly from those three labels (evidence for a hallucinated row,
    faithfulness only when the row shows context, approach only when not ABSTAIN)."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["anotasi"]
    c = {name: i for i, name in enumerate(at.COLUMNS)}
    for row in ws.iter_rows(min_row=2):
        a, h, s = labels.get(row[0].value, ("N", "N", None))
        row[c["abstain"]].value, row[c["hallucinated"]].value, row[c["severity"]].value = a, h, s
        if (h or "").upper() == "Y":
            row[c["wrong_claims"]].value = "fabricated | foo() exists | https://docs.example.org/api"
        if (a or "").upper() == "N":
            row[c["approach_match"]].value = "same"
            if row[c["context"]].value:
                row[c["faithfulness_level"]].value = "all_supported"
        row[c["domain_familiarity"]].value, row[c["uncertain"]].value = "familiar", "N"
        if edit and row[0].value == edit:
            row[3].value = str(row[3].value) + " (edited)"
    wb.save(path)


def test_import_validation_reports_invalid_cells_and_edits(tmp_path, monkeypatch):
    export, key, key_dir, pdir = _export(tmp_path, monkeypatch)
    _fill(pdir / "annotator1.xlsx", {
        "CAL-001": ("Y", None, None), "CAL-002": ("n", "y", "sebagian"),          # valid (case-insensitive)
        "CAL-003": ("X", None, None), "CAL-004": ("Y", "N", None), "CAL-005": ("N", None, None),
        "CAL-006": ("N", "Y", None), "CAL-007": ("N", "N", "PENUH"), "CAL-008": ("N", "Y", "BESAR"),
    }, edit="CAL-009")
    rows, problems = at.load_annotator("calibration", "annotator1", key)
    assert rows["CAL-001"]["category"] == "ABSTAIN" and rows["CAL-002"]["category"] == "HALLUCINATED"
    assert rows["CAL-002"]["severity"] == "SEBAGIAN" and rows["CAL-010"]["category"] == "NOT_HALLUCINATED"
    text = "\n".join(problems)
    for needle in ("(CAL-003): abstain is 'X'", "(CAL-004): hallucinated must be empty when abstain = Y",
                   "(CAL-005): hallucinated is None", "(CAL-006): severity is None",
                   "(CAL-007): severity must be empty when hallucinated = N", "(CAL-008): severity is 'BESAR'",
                   "(CAL-009): column 'candidate_answer' was edited"):
        assert needle in text, needle
    assert len(problems) == 7 and all(rows[i]["category"] is None for i in ("CAL-003", "CAL-006", "CAL-008"))
    assert "annotator1.xlsx!G" in text                       # problems name the cell (abstain is column G)


def test_human_kappa_reference_and_judge_join(tmp_path, monkeypatch):
    export, key, key_dir, pdir = _export(tmp_path, monkeypatch)
    overlap = sorted(k["item_id"] for k in key["items"] if k["in_annotator2_subset"])
    a1_labels = {overlap[0]: ("Y", None, None), overlap[1]: ("N", "Y", "PENUH"), overlap[2]: ("N", "Y", "SEBAGIAN")}
    a2_labels = {overlap[0]: ("Y", None, None), overlap[1]: ("N", "N", None), overlap[2]: ("N", "Y", "PENUH")}
    _fill(pdir / "annotator1.xlsx", a1_labels)
    _fill(pdir / "annotator2.xlsx", a2_labels)
    a1, p1 = at.load_annotator("calibration", "annotator1", key)
    a2, p2 = at.load_annotator("calibration", "annotator2", key)
    assert p1 == [] and p2 == [] and len(a1) == 60 and len(a2) == 30
    hh = at.human_human(a1, a2)
    assert hh["n"] == 30 and [d["item_id"] for d in hh["disagreements"]] == [overlap[1]]   # severity-only differences agree

    ref, problems = at.build_reference("calibration", key, a1, a2, {})
    assert problems == [f"item {overlap[1]}: annotators disagree and no adjudicated label was given"]
    at.write_adjudication_xlsx(pdir / "adjudication.xlsx", hh["disagreements"], a1, a2, {})
    from openpyxl import load_workbook
    wb = load_workbook(pdir / "adjudication.xlsx")
    wb["adjudikasi"]["M2"], wb["adjudikasi"]["N2"], wb["adjudikasi"]["O2"] = "N", "Y", "PENUH"
    wb.save(pdir / "adjudication.xlsx")
    adjudicated, p3 = at.read_adjudication_xlsx(pdir / "adjudication.xlsx")
    ref, problems = at.build_reference("calibration", key, a1, a2, adjudicated)
    assert p3 == [] and problems == [] and len(ref["items"]) == 60
    by_id = {it["item_id"]: it for it in ref["items"]}
    assert by_id[overlap[1]] == {"item_id": overlap[1], "category": "HALLUCINATED", "severity": "PENUH", "source": "adjudicated"}
    assert by_id[overlap[0]]["source"] == "agreed"
    assert sum(1 for it in ref["items"] if it["source"] == "annotator1") == 30
    ref.update({"status": "adjudicated", "n_items": 60, "n_adjudicated": 1, "human_human": None})

    # judge records joined through the key by (run_id, question_id); v2 and v2.1 reported separately
    k_by = {k["item_id"]: k for k in key["items"]}

    def judge(item_id, version, judge_id, attempted=True, claims=(), **kw):
        k = k_by[item_id]
        return {"run_id": k["run_id"], "question_id": k["question_id"], "judge_id": judge_id,
                "prompt_version": version, **_rec(attempted, claims), **kw}

    records = []
    for item_id, h in by_id.items():
        claims = [_claim("FABRICATED", "CORE")] if h["category"] == "HALLUCINATED" else []
        records.append(judge(item_id, "judge-v2.1", "primary", attempted=h["category"] != "ABSTAIN", claims=claims))
        records.append(judge(item_id, "judge-v2", "primary", attempted=True, claims=[]))       # never abstains/flags
    records[0] = {**records[0], "parse_error": True}                                           # one unusable v2.1 record
    results = at.judge_vs_human(ref, key, records)
    assert set(results) == {("judge-v2", "primary"), ("judge-v2.1", "primary")}
    v21 = results[("judge-v2.1", "primary")]
    assert v21["n"] == 59 and v21["n_judge_unusable"] == 1 and v21["kappa_3cat"] == 1.0
    assert results[("judge-v2", "primary")]["kappa_3cat"] == 0.0
    md = at.judge_kappa_markdown("calibration", ref, results)
    assert "| judge-v2.1 | primary | 59 | 1.0 | PASS |" in md and "below target" in md
    assert "binary kappa" in md and "3-category confusion matrix" in md


def test_calibration_generate_plan_and_registry(tmp_path):
    import calibration_generate as cg

    runs = cg.calibration_runs()
    assert [r["run_label"] for r in runs] == ["A [dev]", "B-plain [dev]", "B-grounded [dev]"]
    assert all("_dev_" in r["output_path"].name and "--sample-split" in r["args"] for r in runs)
    assert len({r["config_hash"] for r in runs}) == 3 and all(r["config"]["sample_split"] == "dev" for r in runs)
    assert "--no-require-grounding" in runs[1]["args"] and "--require-grounding" in runs[2]["args"]
    plan = cg.build_plan(runs)
    assert plan["items"] == 150 and plan["est_cost_usd"] > 0 and plan["est_time_sec"] > 0
    assert all(row["pilot"]["prompt_version"] == "v3" and row["pilot"]["model"] == "gpt-4o-mini" for row in plan["rows"])
    reg = tmp_path / "calibration_runs.json"
    cg.write_registry(runs, {"git_commit": "abc", "git_tree_clean": True}, reg, "2026-01-01T00:00:00+00:00")
    data = json.loads(reg.read_text())
    assert data["git_commit"] == "abc" and data["n_annotation_items"] == 60 and len(data["runs"]) == 3
    # the registry it writes is exactly what the judge guard reads
    with pytest.raises(_judge_common.CalibrationGuardError):
        _judge_common.assert_calibration_judging_allowed([runs[0]["output_path"]], registry_path=reg,
                                                         reference_path=tmp_path / "none.json")
    # refuses to generate from a dirty tree, before any subprocess
    assert cg.main(["--dry-run"]) == 0


def test_calibration_registry_records_changed_paths_per_run(tmp_path):
    import calibration_generate as cg

    runs = cg.calibration_runs()
    reg = tmp_path / "calibration_runs.json"
    own = cg.own_output_paths(runs, reg)
    assert reg in own and all(r["output_path"] in own for r in runs)
    runs[1]["changed_paths_at_run_start"] = ["?? llm/a_pure_llm/results/a.jsonl"]
    runs[1]["changed_paths_not_own_output"] = []
    cg.write_registry(runs, {"git_commit": "abc", "git_tree_clean": True}, reg, "t")
    data = json.loads(reg.read_text())
    assert data["runs"][1]["changed_paths_at_run_start"] == ["?? llm/a_pure_llm/results/a.jsonl"]
    assert data["runs"][1]["changed_paths_not_own_output"] == []


def test_calibration_registry_records_data_fingerprint(tmp_path):
    import calibration_generate as cg

    assert cg._sha256_ids([3, 1, 2]) != cg._sha256_ids([1, 2, 3])          # order matters
    assert cg._sha256_ids([1, 2, 3]) == cg._sha256_ids(["1", 2.0, 3])
    data = {"questions_parquet": {"path": "/q.parquet", "size_bytes": 1, "mtime_utc": "t"},
            "n_pool_after_token_filter": 1471, "pool_question_ids_sha256": "p", "dev_question_ids_sha256": "d",
            "dev_question_ids": [1, 2]}
    reg = tmp_path / "calibration_runs.json"
    cg.write_registry(cg.calibration_runs(), {"git_commit": "abc", "git_tree_clean": True}, reg, "t", data)
    assert json.loads(reg.read_text())["data"] == data
    out = tmp_path / "o.jsonl"
    out.write_text('{"question_id": 2}\n{"question_id": 1}\n')
    assert cg.output_question_ids(out) == [2, 1] and cg.output_question_ids(tmp_path / "none") == []


def test_run_metadata_reads_calibration_history_and_relative_output_paths(tmp_path, monkeypatch):
    """A/B write output_path relative to their own folder and keep their history git-ignored: the label
    must resolve from the committed calibration history, from any cwd, with an unchanged run_id."""
    import _run_metadata as rm

    entry = {"output_path": "results/condition_a_openai_gpt-4o-mini_n50_seed42_dev_v3_zzz.jsonl",
             "run_started_at": "2026-01-01T00:00:00+00:00", "status": "success", "sample_split": "dev",
             "prompt_version": "v3"}
    cal = tmp_path / "calibration_run_history.jsonl"
    cal.write_text(json.dumps(entry) + "\n")
    monkeypatch.setattr(rm, "CALIBRATION_HISTORY_PATH", cal)
    monkeypatch.setitem(rm.HISTORY_PATHS, "A", tmp_path / "a_pure_llm" / "logs" / "run_history.jsonl")  # absent
    monkeypatch.chdir(tmp_path)                                    # cwd is NOT the condition folder
    target = tmp_path / "a_pure_llm" / "results" / "condition_a_openai_gpt-4o-mini_n50_seed42_dev_v3_zzz.jsonl"
    meta = rm.resolve_run_metadata(str(target))
    assert meta["run_label"] == "A [dev]" and meta["generator_prompt_version"] == "v3"
    assert meta["run_id"] == rm._history_id("A", entry)
    with pytest.raises(rm.RunMetadataNotFoundError):
        rm.resolve_run_metadata(str(target.with_name("condition_a_other.jsonl")))


def test_calibration_run_history_is_copied_verbatim(tmp_path):
    import calibration_generate as cg

    cwd = tmp_path / "a_pure_llm"
    (cwd / "logs").mkdir(parents=True)
    old = {"output_path": "results/x_dev.jsonl", "run_started_at": "2026-01-01T00:00:00", "status": "no_results"}
    new = {"output_path": "results/x_dev.jsonl", "run_started_at": "2026-01-02T00:00:00", "status": "success", "k": 1}
    other = {"output_path": "results/pilot.jsonl", "run_started_at": "2026-01-03T00:00:00", "status": "success"}
    (cwd / "logs" / "run_history.jsonl").write_text("\n".join(json.dumps(x) for x in (old, new, other)) + "\n")
    out = cg.write_calibration_run_history([{"cwd": cwd, "output_path": cwd / "results" / "x_dev.jsonl"}],
                                           tmp_path / "calibration_run_history.jsonl")
    assert out.read_text() == json.dumps(new) + "\n"
    with pytest.raises(RuntimeError):
        cg.write_calibration_run_history([{"cwd": cwd, "output_path": cwd / "results" / "missing.jsonl"}],
                                         tmp_path / "h.jsonl")
