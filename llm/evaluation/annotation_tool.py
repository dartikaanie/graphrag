"""
annotation_tool.py
=====================================
Blind human annotation for judge calibration and confirmation
(docs/DECISION_C_SCORING.md, Amendment 4 + Addendum).

Subcommands
    export              stratified sample -> one blind XLSX per annotator + key (outside the repo)
    validate            check a filled XLSX (invalid cells are reported, nothing is written)
    human-kappa         annotator 1 vs annotator 2 on the overlap + adjudication template
    finalize-reference  build the adjudicated reference file (unlocks the calibration judge guard)
    judge-kappa         judge vs human reference, per judge and judge version (Markdown report)
    export-reannotation 20% of a COMPLETED sheet again, new ids and order, >= 14 days after completion
    intra-kappa         first pass vs re-annotation of the same annotator
    agreement           every rater pair (judges, researcher, re-annotation, annotator 2) with bootstrap CIs

Two purposes share one flow:
    calibration   60 items = 20 per condition (A, B-plain, B-grounded), seed 42; annotator 2 labels a
                  fixed 30-item subset (10 per condition, seed 42).
    confirmation  63 items = 7 per official test run x the 9 runs R1-R9 (Amendment 11), seed 42; annotator 1
                  only. Exported only when the judge version is frozen (Amendment 11 Record filled in) and all
                  nine official runs are given, each covering exactly the frozen test sample.

BLINDING (Amendment 11). The XLSX files contain no condition, run label, model or judge output.
Annotators see exactly what judge-v2.2 sees, built by the same function (judge_prompt_v2_2.
prepare_judge_input): question title + body, accepted answer, the candidate answer blinded with blind-v3
(citation markers removed, RAG meta phrases normalised; judge_blinding.py), and the retrieval context as
[K1]..[K5] (empty for condition A).

SHEET (Amendment 11, single annotator + optional second one). Shown: item_id, question, accepted_answer,
candidate_answer, context. Filled: annotator_id (prefilled with the sheet's annotator), abstain,
hallucinated, severity, wrong_claims (one wrong claim per line: "<contradicted|fabricated> | <claim> |
<evidence: a URL or a quoted passage>"), faithfulness_level (only when the context cell is not empty),
approach_match, domain_familiarity, uncertain, notes. A row with hallucinated = Y and no wrong claim
with evidence is rejected.

KEY. item_id -> (question_id, run, condition, seed) is written OUTSIDE the repository (default
~/graphrag_annotation_keys/, see --key-dir / GRAPHRAG_ANNOTATION_KEY_DIR). Only its SHA-256 is stored in
the export manifest next to the annotation files; every later step verifies the key against that hash.

CARA PAKAI (from llm/evaluation)
----------------------------------------------------------------------------
    python3 annotation_tool.py export --purpose calibration
    python3 annotation_tool.py validate --xlsx annotation/calibration/annotator1.xlsx --purpose calibration
    python3 annotation_tool.py human-kappa --purpose calibration
    python3 annotation_tool.py finalize-reference --purpose calibration
    python3 annotation_tool.py judge-kappa --purpose calibration --judge-output results/jv2_1_*.jsonl
    python3 annotation_tool.py export-reannotation --purpose calibration          # >= 14 days after `validate`
    python3 annotation_tool.py intra-kappa --purpose calibration

SINGLE ANNOTATOR (Amendment 11). An annotator-2 sheet that is still completely empty is treated as "no
second annotator": finalize-reference then uses annotator 1 alone. A partly filled annotator-2 sheet is an
error. A second annotator can be added later by filling annotator2.xlsx -- no code change.
"""

import argparse
import hashlib
import json
import os
import random
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from judge_blinding import BLINDING_VERSION  # noqa: E402
from judge_binary import (  # noqa: E402
    CATEGORIES, CATEGORY_ABSTAIN, CATEGORY_HALLUCINATED, CATEGORY_NOT_HALLUCINATED, FAITHFULNESS_LEVELS,
    agreement_report, category_of, faithfulness_category, kappa_report,
)

log = print

EVAL_DIR = Path(__file__).resolve().parent
REPO_ROOT = EVAL_DIR.parents[1]
ANNOTATION_ROOT = EVAL_DIR / "annotation"
DEFAULT_KEY_DIR = Path.home() / "graphrag_annotation_keys"
DEFAULT_SEED = 42
ACCEPTANCE_KAPPA = 0.6
REANNOTATION_FRACTION = 0.2
REANNOTATION_MIN_DAYS = 14

PURPOSES = {
    # n_per_stratum / n_overlap_per_stratum: fixed per-condition counts (Amendment 4, calibration set).
    "calibration": {"prefix": "CAL", "n_per_stratum": 20, "n_overlap_per_stratum": 10, "n_total": 60},
    # Amendment 11: 7 per official test run x 9 runs (R1-R9); one annotator.
    "confirmation": {"prefix": "CONF", "n_per_stratum": 7, "n_overlap_per_stratum": 0, "n_total": 63},
}

SHOWN_COLUMNS = ["item_id", "question", "accepted_answer", "candidate_answer", "context"]
FILL_COLUMNS = ["annotator_id", "abstain", "hallucinated", "severity", "wrong_claims", "faithfulness_level",
                "approach_match", "domain_familiarity", "uncertain", "notes"]
COLUMNS = SHOWN_COLUMNS + FILL_COLUMNS
TEXT_COLUMNS = ("question", "accepted_answer", "candidate_answer", "context")
LABEL_COLUMNS = ("abstain", "hallucinated", "severity")
YES_NO = ("Y", "N")
SEVERITIES = ("PENUH", "SEBAGIAN")
WRONG_CLAIM_TYPES = ("contradicted", "fabricated")
APPROACH_MATCH = ("same", "different_valid", "different_unassessable")
FAMILIARITY = ("familiar", "not_familiar")
_EVIDENCE_URL = re.compile(r"https?://\S+")
_EVIDENCE_QUOTE = re.compile(r"[\"\u201c\u201d'].{5,}[\"\u201c\u201d']", re.DOTALL)
MAX_CHARS_EACH = 6000  # same cut as judge_prompt_v2_1.build_messages
_ILLEGAL_XLSX_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_CITATION_MARKER = re.compile(r"\[SO-[^\]]*\]")


class AnnotationError(RuntimeError):
    pass


# ---------------------------------------------------------------------------
# Text shown to annotators == text shown to the judge
# ---------------------------------------------------------------------------

def _cut(s: str) -> str:
    s = s or ""
    return s if len(s) <= MAX_CHARS_EACH else s[:MAX_CHARS_EACH] + "\n[...truncated]"


def _xlsx_safe(s: str) -> str:
    return _ILLEGAL_XLSX_CHARS.sub("", s or "")


def shown_texts(title: str, body_html: str, reference_html: str, candidate: str,
                context_items: list | None = None, has_context: bool = False) -> dict:
    """The text cells, identical to the judge-v2.2 input (judge_prompt_v2_2.prepare_judge_input):
    html_to_text for the question body and the accepted answer, blind-v3 for the candidate, the [K1]..[K5]
    context block (empty cell for condition A); cut at 6000 chars, context items at 2500 each."""
    from _judge_common import html_to_text
    from judge_prompt_v2_2 import prepare_judge_input

    prepared = prepare_judge_input(candidate, context_items, has_context)
    body = html_to_text(body_html) if body_html else "[question body not found in parquet]"
    return {
        "question": _xlsx_safe(f"{title}\n\n{_cut(body)}"),
        "accepted_answer": _xlsx_safe(_cut(html_to_text(reference_html))),
        "candidate_answer": _xlsx_safe(_cut(prepared["candidate_blind"])),
        "context": _xlsx_safe(prepared["context_block"]) if has_context else "",
    }


def _letter(column: str) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(COLUMNS.index(column) + 1)


# Excel writes a carriage return inside a cell as the escape "_x000D_" (OOXML ST_Xstring), and openpyxl reads
# it back literally. Sheet text is converted back before hashing, so this round trip is not taken for an edit.
_XLSX_CR_ESCAPE = "_x000D_"


def from_xlsx_text(value):
    """A text cell as read from a saved sheet, with Excel's "_x000D_" escape turned back into "\r"."""
    return value.replace(_XLSX_CR_ESCAPE, "\r") if isinstance(value, str) else value


def text_sha256(text) -> str:
    """Hash of a text cell, whitespace-normalized so a spreadsheet round trip (line endings, trailing
    spaces) does not look like an edit, while any change of wording does."""
    return hashlib.sha256(" ".join(str(text or "").split()).encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# Sampling (pure, deterministic)
# ---------------------------------------------------------------------------

def _rng(seed: int, *parts) -> random.Random:
    return random.Random("|".join([str(seed), *map(str, parts)]))


def allocate_evenly(strata: list[str], n_total: int) -> dict[str, int]:
    """n_total split as evenly as possible over the sorted strata; the first `remainder` get one extra."""
    strata = sorted(strata)
    base, rem = divmod(n_total, len(strata))
    return {s: base + (1 if i < rem else 0) for i, s in enumerate(strata)}


def sample_items(items: list[dict], purpose: str, seed: int = DEFAULT_SEED,
                 n_per_stratum: int | None = None, n_overlap_per_stratum: int | None = None,
                 n_total: int | None = None) -> list[dict]:
    """Stratified sample by `run_label`. Within each stratum the pool is sorted by question_id and drawn
    with a random.Random seeded from (seed, purpose, stratum) -- independent of every other stratum and of
    any judge output. The annotator-2 subset is drawn from each stratum's sample the same way. Returns the
    sampled items in annotator-1 order (shuffled across strata), each with `item_id` and
    `in_annotator2_subset`; item ids are assigned AFTER the shuffle, so they carry no stratum information."""
    spec = PURPOSES[purpose]
    n_per_stratum = spec["n_per_stratum"] if n_per_stratum is None else n_per_stratum
    n_overlap = spec["n_overlap_per_stratum"] if n_overlap_per_stratum is None else n_overlap_per_stratum
    n_total = spec["n_total"] if n_total is None else n_total

    by_stratum: dict[str, list[dict]] = {}
    for it in items:
        by_stratum.setdefault(it["run_label"], []).append(it)
    quota = ({s: n_per_stratum for s in by_stratum} if n_per_stratum is not None
             else allocate_evenly(list(by_stratum), n_total))

    picked: list[dict] = []
    for stratum in sorted(by_stratum):
        pool = sorted(by_stratum[stratum], key=lambda it: it["question_id"])
        if len(pool) < quota[stratum]:
            raise AnnotationError(f"stratum '{stratum}' has {len(pool)} items, needs {quota[stratum]}")
        chosen = _rng(seed, purpose, "sample", stratum).sample(pool, quota[stratum])
        overlap_ids = {id(x) for x in _rng(seed, purpose, "overlap", stratum).sample(
            sorted(chosen, key=lambda it: it["question_id"]), min(n_overlap, len(chosen)))}
        picked += [{**it, "in_annotator2_subset": id(it) in overlap_ids} for it in chosen]

    picked.sort(key=lambda it: (it["run_label"], it["question_id"]))
    _rng(seed, purpose, "order", "annotator1").shuffle(picked)
    width = max(3, len(str(len(picked))))
    for i, it in enumerate(picked, start=1):
        it["item_id"] = f"{spec['prefix']}-{i:0{width}d}"
    return picked


def annotator2_order(sampled: list[dict], purpose: str, seed: int = DEFAULT_SEED) -> list[dict]:
    subset = sorted((it for it in sampled if it["in_annotator2_subset"]), key=lambda it: it["item_id"])
    _rng(seed, purpose, "order", "annotator2").shuffle(subset)
    return subset


# ---------------------------------------------------------------------------
# XLSX writing / reading
# ---------------------------------------------------------------------------

def write_annotation_xlsx(path: Path, rows: list[dict], annotator_id: str = "annotator1") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "anotasi"
    ws.append(COLUMNS)
    for r in rows:
        ws.append([r[c] for c in SHOWN_COLUMNS] + [annotator_id] + [None] * (len(FILL_COLUMNS) - 1))

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    widths = {"item_id": 11, "question": 55, "accepted_answer": 55, "candidate_answer": 55, "context": 55,
              "annotator_id": 13, "wrong_claims": 50, "notes": 32}
    for name in COLUMNS:
        ws.column_dimensions[_letter(name)].width = widths.get(name, 16)
    for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    last = len(rows) + 1
    for name, values in (("abstain", YES_NO), ("hallucinated", YES_NO), ("severity", SEVERITIES),
                         ("faithfulness_level", FAITHFULNESS_LEVELS), ("approach_match", APPROACH_MATCH),
                         ("domain_familiarity", FAMILIARITY), ("uncertain", YES_NO)):
        col = _letter(name)
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True, showErrorMessage=True)
        dv.error, dv.errorTitle = f"Isi dengan salah satu: {', '.join(values)}", "Nilai tidak valid"
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")

    info = wb.create_sheet("petunjuk")
    for line in (
        "Baca docs/ANNOTATION_GUIDELINE.md sebelum mulai.",
        "Panduan singkat: docs/ANNOTATION_QUICKSTART.md.",
        "annotator_id: sudah terisi; jangan diubah kecuali Anda annotator lain.",
        "abstain: Y jika jawaban kandidat tidak mencoba menjawab; N jika mencoba menjawab.",
        "hallucinated: isi HANYA jika abstain = N. Y jika ada >= 1 klaim yang bertentangan dengan referensi "
        "atau fakta mapan, atau yang dikarang (API/fungsi/opsi tidak ada). N jika tidak ada.",
        "severity: isi HANYA jika hallucinated = Y. PENUH jika kesalahan mengenai solusi utama; SEBAGIAN jika "
        "hanya detail sekunder.",
        "wrong_claims: WAJIB jika hallucinated = Y, kosong jika tidak. Satu klaim salah per baris (Alt+Enter): "
        "jenis | klaim | bukti. jenis = contradicted atau fabricated. bukti = URL dokumentasi, atau kutipan "
        "dalam tanda petik dari thread/konteks (mis. \"[K2] ... teks ...\").",
        "faithfulness_level: isi HANYA jika kolom context tidak kosong dan abstain = N: all_supported / "
        "partly_supported / not_supported (klaim jawaban didukung konteks?).",
        "approach_match: isi jika abstain = N: same / different_valid / different_unassessable.",
        "domain_familiarity: selalu isi: familiar / not_familiar (keakraban Anda dengan topiknya).",
        "uncertain: selalu isi: Y jika Anda ragu pada label baris ini, N jika tidak.",
        "notes: catatan bebas. Tulis 'referensi usang' bila memakai aturan referensi usang.",
        "Jangan mengubah kolom item_id, question, accepted_answer, candidate_answer, context.",
    ):
        info.append([line])
    info.column_dimensions["A"].width = 120
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _norm_label(value) -> str | None:
    if value is None:
        return None
    v = str(value).strip().upper()
    return v or None


def _norm_choice(value) -> str | None:
    if value is None:
        return None
    v = str(value).strip().lower()
    return v or None


def parse_wrong_claims(cell) -> tuple[list[dict], list[str]]:
    """One wrong claim per non-empty line: "<contradicted|fabricated> | <claim> | <evidence>". Evidence is a
    URL or a quoted passage (thread or context). Returns (claims, problems)."""
    claims, problems = [], []
    for n, line in enumerate(str(cell or "").splitlines(), start=1):
        if not line.strip():
            continue
        parts = [x.strip() for x in line.split("|", 2)]
        if len(parts) != 3:
            problems.append(f"line {n}: expected 'type | claim | evidence', got {line.strip()[:60]!r}")
            continue
        kind, claim, evidence = parts[0].lower(), parts[1], parts[2]
        if kind not in WRONG_CLAIM_TYPES:
            problems.append(f"line {n}: type is {parts[0]!r}, must be contradicted or fabricated")
        if not claim:
            problems.append(f"line {n}: the claim is empty")
        if not (_EVIDENCE_URL.search(evidence) or _EVIDENCE_QUOTE.search(evidence)):
            problems.append(f"line {n}: evidence must be a URL or a quoted passage, got {evidence[:60]!r}")
        claims.append({"type": kind, "claim": claim, "evidence": evidence})
    return claims, problems


def read_annotation_xlsx(path: Path) -> tuple[list[dict], list[str]]:
    """Reads a filled annotation sheet. Returns (rows, problems); every problem names its cell. Rules:
      annotator_id        required
      abstain             required, Y/N
      hallucinated        required iff abstain == N, Y/N; empty when abstain == Y
      severity            required iff hallucinated == Y, PENUH/SEBAGIAN; empty otherwise
      wrong_claims        required iff hallucinated == Y: >= 1 line "type | claim | evidence" with evidence
                          (URL or quoted passage); empty otherwise
      faithfulness_level  required iff abstain == N and the context cell is not empty; empty otherwise
      approach_match      required iff abstain == N; empty when abstain == Y
      domain_familiarity  required; uncertain required (Y/N)
    A row with any problem has category None."""
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)["anotasi"]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    problems: list[str] = []
    if header[:len(COLUMNS)] != COLUMNS:
        return [], [f"{path.name}: header row is {header}, expected {COLUMNS} (a sheet from an older export "
                    f"cannot be read; re-export with the current format)"]
    col = {name: i for i, name in enumerate(COLUMNS)}

    rows, seen = [], set()
    for n, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        cells = list(cells) + [None] * (len(COLUMNS) - len(cells))
        if all(v is None or str(v).strip() == "" for v in cells[:len(COLUMNS)]):
            continue
        item_id = str(cells[col["item_id"]] or "").strip()
        if not item_id:
            problems.append(f"{path.name}!A{n}: missing item_id")
            continue
        if item_id in seen:
            problems.append(f"{path.name}!A{n}: duplicate item_id {item_id}")
        seen.add(item_id)

        def cell(name):
            return cells[col[name]]

        def bad(name, msg):
            row_problems.append(f"{path.name}!{_letter(name)}{n} ({item_id}): {msg}")

        row_problems: list[str] = []
        annotator_id = str(cell("annotator_id") or "").strip()
        abstain, hall, sev = (_norm_label(cell(c)) for c in LABEL_COLUMNS)
        faith, approach, familiarity = (_norm_choice(cell(c)) for c in ("faithfulness_level", "approach_match",
                                                                          "domain_familiarity"))
        uncertain = _norm_label(cell("uncertain"))
        has_context = bool(str(cell("context") or "").strip())
        wrong_claims, wc_problems = parse_wrong_claims(cell("wrong_claims"))

        if not annotator_id:
            bad("annotator_id", "annotator_id is empty")
        if abstain not in YES_NO:
            bad("abstain", f"abstain is {abstain!r}, must be Y or N")
        elif abstain == "Y":
            for name, value in (("hallucinated", hall), ("severity", sev), ("faithfulness_level", faith),
                                ("approach_match", approach)):
                if value is not None:
                    bad(name, f"{name} must be empty when abstain = Y (is {value!r})")
            if wrong_claims:
                bad("wrong_claims", "wrong_claims must be empty when abstain = Y")
        else:
            if hall not in YES_NO:
                bad("hallucinated", f"hallucinated is {hall!r}, must be Y or N when abstain = N")
            elif hall == "Y":
                if sev not in SEVERITIES:
                    bad("severity", f"severity is {sev!r}, must be PENUH or SEBAGIAN when hallucinated = Y")
                if not wrong_claims:
                    bad("wrong_claims", "hallucinated = Y needs at least one wrong claim with evidence "
                                        "('type | claim | evidence')")
                for wp in wc_problems:
                    bad("wrong_claims", wp)
            else:
                if sev is not None:
                    bad("severity", f"severity must be empty when hallucinated = N (is {sev!r})")
                if wrong_claims:
                    bad("wrong_claims", "wrong_claims must be empty when hallucinated = N")
            if has_context and faith not in FAITHFULNESS_LEVELS:
                bad("faithfulness_level", f"faithfulness_level is {faith!r}, must be one of {FAITHFULNESS_LEVELS} "
                                          f"(the context cell is not empty)")
            if not has_context and faith is not None:
                bad("faithfulness_level", "faithfulness_level must be empty when the context cell is empty")
            if approach not in APPROACH_MATCH:
                bad("approach_match", f"approach_match is {approach!r}, must be one of {APPROACH_MATCH}")
        if familiarity not in FAMILIARITY:
            bad("domain_familiarity", f"domain_familiarity is {familiarity!r}, must be one of {FAMILIARITY}")
        if uncertain not in YES_NO:
            bad("uncertain", f"uncertain is {uncertain!r}, must be Y or N")
        problems += row_problems

        category = None
        if not row_problems:
            category = (CATEGORY_ABSTAIN if abstain == "Y"
                        else CATEGORY_HALLUCINATED if hall == "Y" else CATEGORY_NOT_HALLUCINATED)
        rows.append({
            "item_id": item_id, "row": n, "annotator_id": annotator_id, "abstain": abstain, "hallucinated": hall,
            "severity": sev, "wrong_claims": wrong_claims, "faithfulness_level": faith, "approach_match": approach,
            "domain_familiarity": familiarity, "uncertain": uncertain, "has_context": has_context,
            "notes": str(cell("notes") or "").strip(), "category": category,
            "text_sha256": {c: text_sha256(from_xlsx_text(cell(c))) for c in TEXT_COLUMNS},
        })
    return rows, problems


# ---------------------------------------------------------------------------
# Key (outside the repo) + manifest (inside, hash only)
# ---------------------------------------------------------------------------

def key_dir_from(arg: str | None) -> Path:
    return Path(arg or os.getenv("GRAPHRAG_ANNOTATION_KEY_DIR") or DEFAULT_KEY_DIR).expanduser()


def purpose_dir(purpose: str, annotation_root: Path | None = None) -> Path:
    return (annotation_root or ANNOTATION_ROOT) / purpose


def file_sha256(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def load_manifest(purpose: str, annotation_root: Path | None = None) -> dict:
    path = purpose_dir(purpose, annotation_root) / "export_manifest.json"
    if not path.exists():
        raise AnnotationError(f"no export manifest at {path} -- run `export --purpose {purpose}` first")
    return json.loads(path.read_text())


def load_verified_key(purpose: str, key_dir: Path, annotation_root: Path | None = None) -> dict:
    """Loads the key from OUTSIDE the repo and verifies its SHA-256 against the committed manifest."""
    manifest = load_manifest(purpose, annotation_root)
    key_path = key_dir / manifest["key_file_name"]
    if not key_path.exists():
        raise AnnotationError(f"key file not found at {key_path} (set --key-dir or GRAPHRAG_ANNOTATION_KEY_DIR)")
    actual = file_sha256(key_path)
    if actual != manifest["key_sha256"]:
        raise AnnotationError(f"key file {key_path} has sha256 {actual}, but the export manifest records "
                              f"{manifest['key_sha256']} -- this is not the key of this export")
    return json.loads(key_path.read_text())


def check_against_key(rows: list[dict], key: dict, annotator: str, file_name: str) -> list[str]:
    """Item-id coverage and text integrity of a filled sheet against the key."""
    problems = []
    expected = {it["item_id"]: it for it in key["items"] if annotator == "annotator1" or it["in_annotator2_subset"]}
    got = {r["item_id"] for r in rows}
    for missing in sorted(set(expected) - got):
        problems.append(f"{file_name}: item {missing} is missing")
    for extra in sorted(got - set(expected)):
        problems.append(f"{file_name}: unexpected item {extra}")
    for r in rows:
        k = expected.get(r["item_id"])
        if k is None:
            continue
        for c in TEXT_COLUMNS:
            if r["text_sha256"][c] != k["text_sha256"][c]:
                problems.append(f"{file_name} row {r['row']} ({r['item_id']}): column '{c}' was edited "
                                f"(text no longer matches the export)")
    return problems


# ---------------------------------------------------------------------------
# export
# ---------------------------------------------------------------------------

def load_run_items(run_files: list[str]) -> list[dict]:
    from _run_metadata import resolve_run_metadata

    items = []
    for path in run_files:
        meta = resolve_run_metadata(path)
        for line in open(path):
            if not line.strip():
                continue
            rec = json.loads(line)
            items.append({
                "run_id": meta["run_id"], "run_label": meta["run_label"], "condition": meta["condition"],
                "source_file": str(path), "config_hash": rec.get("config_hash"),
                "question_id": rec["question_id"], "title": rec.get("title", "") or "",
                "ground_truth_answer_html": rec.get("ground_truth_answer", "") or "",
                "candidate": rec.get("llm_answer", "") or "",
                "retrieved_context": rec.get("retrieved_context") or [],
                "has_context": str(meta["condition"]).upper() != "A",
            })
    return items


def build_export(items: list[dict], bodies_html: dict, purpose: str, seed: int = DEFAULT_SEED, **sample_kw) -> dict:
    """Pure: returns {"annotator1": [rows], "annotator2": [rows], "key_items": [...], "n_markers_stripped"}.
    Rows hold ONLY the blind columns."""
    sampled = sample_items(items, purpose, seed, **sample_kw)
    rows_by_id, key_items, n_markers = {}, [], 0
    for it in sampled:
        has_context = it.get("has_context", str(it.get("condition", "")).upper() != "A")
        texts = shown_texts(it["title"], bodies_html.get(it["question_id"], ""), it["ground_truth_answer_html"],
                            it["candidate"], it.get("retrieved_context"), has_context)
        markers = len(_CITATION_MARKER.findall(it["candidate"]))
        n_markers += markers
        rows_by_id[it["item_id"]] = {"item_id": it["item_id"], **texts}
        key_items.append({
            "item_id": it["item_id"], "question_id": it["question_id"], "run_id": it["run_id"],
            "run_label": it["run_label"], "condition": it["condition"], "config_hash": it["config_hash"],
            "source_file": Path(it["source_file"]).name, "in_annotator2_subset": it["in_annotator2_subset"],
            "citation_markers_stripped": markers, "has_context": has_context,
            "text_sha256": {c: text_sha256(texts[c]) for c in TEXT_COLUMNS},
        })
    return {
        "annotator1": [rows_by_id[it["item_id"]] for it in sampled],
        "annotator2": [rows_by_id[it["item_id"]] for it in annotator2_order(sampled, purpose, seed)],
        "key_items": key_items, "n_markers_stripped": n_markers,
    }


DECISION_RECORD = REPO_ROOT / "docs" / "DECISION_C_SCORING.md"


def frozen_judge_version(decision_path: Path = DECISION_RECORD) -> str | None:
    """The judge version in Amendment 11's Record ("- Judge version frozen: judge-v2.x  commit: ..."), or None
    while it is still "[ ]"."""
    text = Path(decision_path).read_text()
    start = text.find("## Amendment 11")
    record = text.find("### Record (fill in)", start)
    if start < 0 or record < 0:
        return None
    for line in text[record:].splitlines()[1:12]:
        if line.startswith("- Judge version frozen:"):
            m = re.search(r"judge-v\d+(?:\.\d+)*", line.split("commit:")[0])
            return m.group(0) if m else None
    return None


def official_cell(config: dict, manifest: dict | None, official: dict | None) -> str | None:
    """Official run-matrix label (R1-R9, llm/official_run_matrix.py) of a test-sample run, or None."""
    cond = str(config.get("condition") or "").upper()
    grounded = "grounded" if config.get("require_grounding") else "plain"
    if cond == "A":
        return "A"
    if cond in ("B", "D"):
        return f"{cond}-{grounded}"
    if cond == "C":
        from llm.c_graphrag.official_params import role_for_alpha

        role = (manifest or {}).get("official_role") or role_for_alpha(official, config.get("alpha"))
        return {"c_primary": f"C-primary-{grounded}", "c_ablation": f"C-ablation-{grounded}"}.get(role)
    return None


def confirmation_cells(run_files: list[str], test_ids: set, official: dict | None) -> dict[str, str]:
    """{run file: cell label} for exactly the nine official runs; AnnotationError otherwise."""
    from llm.official_run_matrix import CELLS

    expected = {label for _, label, _, _, _ in CELLS}
    cells, problems = {}, []
    for path in run_files:
        recs = [json.loads(line) for line in open(path) if line.strip()]
        cfg = (recs[0].get("config") if recs else None) or {}
        man_path = Path(path).with_name(Path(path).name + ".manifest.json")
        manifest = json.loads(man_path.read_text()) if man_path.exists() else None
        if cfg.get("sample_split") != "test":
            problems.append(f"{Path(path).name}: sample_split is {cfg.get('sample_split')!r}, not 'test'")
            continue
        ids = {int(r["question_id"]) for r in recs}
        if ids != test_ids:
            problems.append(f"{Path(path).name}: does not cover exactly the frozen test sample "
                            f"({len(ids - test_ids)} extra, {len(test_ids - ids)} missing)")
            continue
        cell = official_cell(cfg, manifest, official)
        if cell not in expected:
            problems.append(f"{Path(path).name}: not one of the official runs R1-R9 (cell {cell!r})")
        elif cell in cells.values():
            problems.append(f"{Path(path).name}: a second run for {cell}")
        else:
            cells[str(path)] = cell
    missing = sorted(expected - set(cells.values()))
    if missing:
        problems.append(f"official runs missing: {missing}")
    if problems:
        raise AnnotationError("confirmation export refused: " + "; ".join(problems))
    return cells


def check_confirmation_allowed(run_files: list[str], decision_path: Path = DECISION_RECORD,
                               test_ids: set | None = None, official: dict | None = None) -> dict:
    """Amendment 11: refuse the confirmation export until the judge is frozen and all nine official runs exist."""
    version = frozen_judge_version(decision_path)
    if version is None:
        raise AnnotationError("confirmation export refused: the judge version is not frozen yet (the Record of "
                              "Amendment 11 in docs/DECISION_C_SCORING.md is still empty)")
    if test_ids is None:
        from llm.sample_split import load_frozen_ids
        test_ids = {int(q) for q in load_frozen_ids("test")["ids"]}
    if official is None:
        from llm.c_graphrag.official_params import load_official_params
        official = load_official_params()
    return {"judge_version": version, "cells": confirmation_cells(run_files, test_ids, official)}


def cmd_export(args) -> int:
    from _judge_common import load_calibration_registry, load_question_bodies
    from llm.repo_state import repo_state

    purpose, out_dir, key_dir = args.purpose, purpose_dir(args.purpose), key_dir_from(args.key_dir)
    run_files = args.run_files
    if not run_files and purpose == "calibration":
        run_files = [str(REPO_ROOT / r["output_path"]) for r in load_calibration_registry().get("runs", [])]
    if not run_files:
        raise AnnotationError("--run-files is required (no calibration registry found)")
    manifest_path = out_dir / "export_manifest.json"
    if manifest_path.exists() and not args.force:
        raise AnnotationError(f"{manifest_path} already exists -- an export for '{purpose}' was already made. "
                              f"Re-exporting would invalidate annotation in progress; pass --force to overwrite.")
    questions_parquet = args.questions_parquet or os.getenv("QUESTIONS_PARQUET")
    if not questions_parquet:
        raise AnnotationError("--questions-parquet (or QUESTIONS_PARQUET in .env) is required")

    confirmation = None
    if purpose == "confirmation":
        confirmation = check_confirmation_allowed(run_files)
    items = load_run_items(run_files)
    if confirmation:  # strata = the nine official cells, not the run labels
        for it in items:
            it["run_label"] = confirmation["cells"][str(it["source_file"])]
    sample_kw = {}
    if args.n_total is not None:
        sample_kw["n_total"] = args.n_total
    sampled_preview = sample_items(items, purpose, args.seed, **sample_kw)
    bodies = load_question_bodies(questions_parquet, sorted({it["question_id"] for it in sampled_preview}))
    export = build_export(items, bodies, purpose, args.seed, **sample_kw)

    spec, git = PURPOSES[purpose], repo_state()
    key = {
        "purpose": purpose, "seed": args.seed, "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "git_commit": git["git_commit"],
        "sampling": {"strata": "run_label", "n_per_stratum": spec["n_per_stratum"],
                     "n_overlap_per_stratum": spec["n_overlap_per_stratum"],
                     "n_total": len(export["key_items"]), "run_files": [Path(f).name for f in run_files]},
        "items": export["key_items"],
        **({"frozen_judge_version": confirmation["judge_version"]} if confirmation else {}),
    }
    key_dir.mkdir(parents=True, exist_ok=True)
    key_name = f"{purpose}_key_seed{args.seed}.json"
    key_path = key_dir / key_name
    key_path.write_text(json.dumps(key, indent=2) + "\n")

    files = {"annotator1": out_dir / "annotator1.xlsx"}
    write_annotation_xlsx(files["annotator1"], export["annotator1"], "annotator1")
    if export["annotator2"]:
        files["annotator2"] = out_dir / "annotator2.xlsx"
        write_annotation_xlsx(files["annotator2"], export["annotator2"], "annotator2")

    manifest_path.write_text(json.dumps({
        "purpose": purpose, "seed": args.seed, "created_at_utc": key["created_at_utc"],
        "git_commit": git["git_commit"],
        "n_items": len(export["annotator1"]), "n_items_annotator2": len(export["annotator2"]),
        "columns": COLUMNS,
        "key_file_name": key_name, "key_sha256": file_sha256(key_path),
        "key_location_note": "kept OUTSIDE the repository until adjudication is complete (Amendment 4 Addendum)",
        "blinding": {
            "blinding_version": BLINDING_VERSION,
            "hidden": ["condition", "run_label", "run_id", "model", "judge output"],
            "citation_markers_stripped": True, "n_citation_markers_stripped": export["n_markers_stripped"],
            "candidate": "judge_prompt_v2_2.prepare_judge_input: blind-v3 (citation markers removed, RAG meta "
                         "phrases normalised; judge_blinding.py), identical to the judge-v2.2 input",
            "context": "[K1]..[K5] block as in the judge-v2.2 input (2500 chars per item); empty for condition A",
            "text_conversion": "html_to_text for question body and accepted answer; cut at 6000 chars, as the judge input",
        },
    }, indent=2) + "\n")

    log(f"Wrote {len(export['annotator1'])} items -> {files['annotator1']}")
    if "annotator2" in files:
        log(f"Wrote {len(export['annotator2'])} items -> {files['annotator2']}")
    log(f"Wrote key OUTSIDE the repo -> {key_path}")
    log(f"Wrote manifest (key sha256 only) -> {manifest_path}")
    return 0


# ---------------------------------------------------------------------------
# validate / human-kappa / finalize-reference
# ---------------------------------------------------------------------------

def load_annotator(purpose: str, annotator: str, key: dict, annotation_root: Path | None = None,
                   path: Path | None = None) -> tuple[dict, list[str]]:
    path = path or purpose_dir(purpose, annotation_root) / f"{annotator}.xlsx"
    if not path.exists():
        return {}, [f"{path} not found"]
    rows, problems = read_annotation_xlsx(path)
    problems += check_against_key(rows, key, annotator, path.name)
    return {r["item_id"]: r for r in rows}, problems


def completion_path(purpose: str, annotator: str, annotation_root: Path | None = None) -> Path:
    return purpose_dir(purpose, annotation_root) / f"{annotator}_completed.json"


def record_completion(purpose: str, annotator: str, sheet: Path, now: datetime | None = None,
                      annotation_root: Path | None = None) -> dict:
    """Writes <annotator>_completed.json the first time a sheet validates completely (or again when the sheet
    changed since). Its completed_at_utc starts the re-annotation waiting period."""
    path = completion_path(purpose, annotator, annotation_root)
    sha = file_sha256(sheet)
    if path.exists():
        rec = json.loads(path.read_text())
        if rec.get("sheet_sha256") == sha:
            return rec
    rec = {"purpose": purpose, "annotator": annotator, "sheet": Path(sheet).name, "sheet_sha256": sha,
           "completed_at_utc": (now or datetime.now(timezone.utc)).isoformat()}
    path.write_text(json.dumps(rec, indent=2) + "\n")
    return rec


def cmd_validate(args) -> int:
    key = load_verified_key(args.purpose, key_dir_from(args.key_dir))
    annotator = args.annotator or ("annotator2" if "annotator2" in Path(args.xlsx).name else "annotator1")
    rows, problems = load_annotator(args.purpose, annotator, key, path=Path(args.xlsx))
    log(f"{args.xlsx}: {len(rows)} rows read as {annotator}; {len(problems)} problem(s)")
    for p in problems:
        log(f"  - {p}")
    if not problems and Path(args.xlsx).resolve() == (purpose_dir(args.purpose) / f"{annotator}.xlsx").resolve():
        rec = record_completion(args.purpose, annotator, Path(args.xlsx))
        log(f"Sheet complete; completion recorded at {rec['completed_at_utc']} "
            f"(re-annotation possible after {REANNOTATION_MIN_DAYS} days)")
    return 1 if problems else 0


def sheet_started(path: Path) -> bool:
    """True when any annotator-filled cell (other than the prefilled annotator_id) is non-empty."""
    from openpyxl import load_workbook

    if not Path(path).exists():
        return False
    ws = load_workbook(path, data_only=True)["anotasi"]
    filled = [COLUMNS.index(c) for c in FILL_COLUMNS if c != "annotator_id"]
    return any(cells[i] not in (None, "") for cells in ws.iter_rows(min_row=2, values_only=True)
               for i in filled if i < len(cells))


def human_human(a1: dict, a2: dict) -> dict:
    """Pre-adjudication agreement on the overlap items + the disagreement list."""
    overlap = sorted(set(a1) & set(a2))
    report = agreement_report([a1[i]["category"] for i in overlap], [a2[i]["category"] for i in overlap])
    report["item_ids"] = overlap
    report["disagreements"] = [
        {"item_id": i, "annotator1": a1[i]["category"], "annotator2": a2[i]["category"],
         "annotator1_severity": a1[i]["severity"], "annotator2_severity": a2[i]["severity"]}
        for i in overlap if a1[i]["category"] != a2[i]["category"]
    ]
    return report


def write_adjudication_xlsx(path: Path, disagreements: list[dict], a1: dict, a2: dict, sheet_rows: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "adjudikasi"
    header = ["item_id", "question", "accepted_answer", "candidate_answer",
              "a1_abstain", "a1_hallucinated", "a1_severity", "a1_notes",
              "a2_abstain", "a2_hallucinated", "a2_severity", "a2_notes",
              "abstain", "hallucinated", "severity", "notes"]
    ws.append(header)
    for d in disagreements:
        i = d["item_id"]
        t = sheet_rows.get(i, {})
        ws.append([i, t.get("question"), t.get("accepted_answer"), t.get("candidate_answer"),
                   a1[i]["abstain"], a1[i]["hallucinated"], a1[i]["severity"], a1[i]["notes"],
                   a2[i]["abstain"], a2[i]["hallucinated"], a2[i]["severity"], a2[i]["notes"],
                   None, None, None, None])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "B2"
    for row in ws.iter_rows(min_row=1, max_row=len(disagreements) + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in zip("ABCD", (11, 50, 50, 50)):
        ws.column_dimensions[col].width = width
    last = len(disagreements) + 1
    for col, values in (("M", YES_NO), ("N", YES_NO), ("O", SEVERITIES)):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max(last, 2)}")
    wb.save(path)


def read_adjudication_xlsx(path: Path) -> tuple[dict, list[str]]:
    """Final labels of the disagreement items: same validation rules as an annotation sheet, on the
    final abstain/hallucinated/severity columns (M, N, O)."""
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)["adjudikasi"]
    out, problems = {}, []
    for n, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not cells or not cells[0]:
            continue
        item_id = str(cells[0]).strip()
        abstain, hall, sev = (_norm_label(v) for v in cells[12:15])
        ok = (abstain == "Y" and hall is None and sev is None) or \
             (abstain == "N" and hall == "N" and sev is None) or \
             (abstain == "N" and hall == "Y" and sev in SEVERITIES)
        if not ok:
            problems.append(f"{path.name} row {n} ({item_id}): final labels abstain={abstain!r} "
                            f"hallucinated={hall!r} severity={sev!r} are missing or inconsistent (columns M-O)")
            continue
        out[item_id] = {
            "abstain": abstain, "hallucinated": hall, "severity": sev,
            "category": CATEGORY_ABSTAIN if abstain == "Y" else
            CATEGORY_HALLUCINATED if hall == "Y" else CATEGORY_NOT_HALLUCINATED,
        }
    return out, problems


def _texts_from_sheet(path: Path) -> dict:
    from openpyxl import load_workbook

    out = {}
    for cells in load_workbook(path, data_only=True)["anotasi"].iter_rows(min_row=2, values_only=True):
        if cells and cells[0]:
            out[str(cells[0]).strip()] = {"question": cells[1], "accepted_answer": cells[2],
                                          "candidate_answer": cells[3], "context": cells[4] if len(cells) > 4 else ""}
    return out


def cmd_human_kappa(args) -> int:
    purpose, pdir = args.purpose, purpose_dir(args.purpose)
    key = load_verified_key(purpose, key_dir_from(args.key_dir))
    a1, p1 = load_annotator(purpose, "annotator1", key)
    a2, p2 = load_annotator(purpose, "annotator2", key)
    problems = p1 + p2
    if problems:
        log(f"[STOP] {len(problems)} problem(s) in the annotation files -- fix them first:")
        for p in problems:
            log(f"  - {p}")
        return 1
    report = human_human(a1, a2)
    adjudication_path = pdir / "adjudication.xlsx"
    if not adjudication_path.exists() or args.force:
        write_adjudication_xlsx(adjudication_path, report["disagreements"], a1, a2,
                                _texts_from_sheet(pdir / "annotator1.xlsx"))
    (pdir / "human_human_agreement.json").write_text(json.dumps(report, indent=2) + "\n")
    log(f"Human-human agreement on {report['n']} overlap items (pre-adjudication):")
    log(f"  3-category kappa (unweighted): {report['kappa_3cat']}  observed agreement: {report['observed_agreement_3cat']}")
    log(f"  binary kappa on both-non-ABSTAIN (n={report['n_both_non_abstain']}): {report['kappa_binary_non_abstain']}")
    log(f"  disagreements: {len(report['disagreements'])} -> fill the final columns in {adjudication_path}")
    return 0


def build_reference(purpose: str, key: dict, a1: dict, a2: dict, adjudicated: dict) -> tuple[dict, list[str]]:
    """Reference label per item: annotator 1's label; on overlap items where the two annotators disagree,
    the adjudicated label (required). Returns (reference, problems)."""
    problems, items = [], []
    for k in key["items"]:
        i = k["item_id"]
        if i not in a1:
            problems.append(f"item {i}: no annotator-1 label")
            continue
        label, source = a1[i], "annotator1"
        if i in a2:
            if a2[i]["category"] == a1[i]["category"]:
                source = "agreed"
            elif i in adjudicated:
                label, source = adjudicated[i], "adjudicated"
            else:
                problems.append(f"item {i}: annotators disagree and no adjudicated label was given")
                continue
        items.append({"item_id": i, "category": label["category"], "severity": label.get("severity"),
                      "source": source})
    return {"purpose": purpose, "items": items}, problems


def cmd_finalize_reference(args) -> int:
    purpose, pdir = args.purpose, purpose_dir(args.purpose)
    manifest = load_manifest(purpose)
    key = load_verified_key(purpose, key_dir_from(args.key_dir))
    a1, problems = load_annotator(purpose, "annotator1", key)
    a2, adjudicated, hh = {}, {}, None
    if manifest.get("n_items_annotator2") and not sheet_started(pdir / "annotator2.xlsx"):
        log("annotator2.xlsx is empty: single-annotator reference (Amendment 11)")
    elif manifest.get("n_items_annotator2"):
        a2, p2 = load_annotator(purpose, "annotator2", key)
        problems += p2
        if not problems:
            hh = human_human(a1, a2)
            if hh["disagreements"]:
                adj_path = pdir / "adjudication.xlsx"
                if not adj_path.exists():
                    problems.append(f"{adj_path} not found -- run `human-kappa` and fill it")
                else:
                    adjudicated, p3 = read_adjudication_xlsx(adj_path)
                    problems += p3
    reference, p4 = ({}, []) if problems else build_reference(purpose, key, a1, a2, adjudicated)
    problems += p4
    if problems:
        log(f"[STOP] cannot finalize the reference -- {len(problems)} problem(s):")
        for p in problems:
            log(f"  - {p}")
        return 1
    reference.update({
        "status": "adjudicated", "finalized_at_utc": datetime.now(timezone.utc).isoformat(),
        "key_sha256": manifest["key_sha256"], "n_items": len(reference["items"]),
        "human_human": None if hh is None else {k: hh[k] for k in (
            "n", "kappa_3cat", "observed_agreement_3cat", "n_both_non_abstain", "kappa_binary_non_abstain")},
        "n_adjudicated": sum(1 for it in reference["items"] if it["source"] == "adjudicated"),
    })
    out = pdir / "reference_adjudicated.json"
    out.write_text(json.dumps(reference, indent=2) + "\n")
    log(f"Reference with {reference['n_items']} items -> {out}")
    return 0


# ---------------------------------------------------------------------------
# judge-kappa
# ---------------------------------------------------------------------------

def judge_vs_human(reference: dict, key: dict, judge_records: list[dict]) -> dict:
    """{(prompt_version, judge_id): agreement} -- 3-category kappa on all reference items with a usable
    judge record, binary kappa on items BOTH label non-ABSTAIN. Items whose judge record is missing or
    unusable (parse error / failed call) are excluded and counted, never guessed."""
    human = {it["item_id"]: it["category"] for it in reference["items"]}
    key_by_item = {k["item_id"]: k for k in key["items"]}
    index: dict[tuple, dict] = {}
    for r in judge_records:
        index[(r.get("prompt_version"), r["judge_id"], r["run_id"], r["question_id"])] = r

    out = {}
    for version, judge_id in sorted({(k[0], k[1]) for k in index}, key=str):
        h_labels, j_labels, n_missing, n_unusable = [], [], 0, 0
        for item_id, h in sorted(human.items()):
            k = key_by_item[item_id]
            rec = index.get((version, judge_id, k["run_id"], k["question_id"]))
            if rec is None:
                n_missing += 1
                continue
            cat = category_of(rec)
            if cat is None:
                n_unusable += 1
                continue
            h_labels.append(h)
            j_labels.append(cat)
        rep = agreement_report(h_labels, j_labels) if h_labels else {"n": 0, "kappa_3cat": None}
        rep.update({"n_reference_items": len(human), "n_judge_missing": n_missing, "n_judge_unusable": n_unusable})
        out[(version, judge_id)] = rep
    return out


def _matrix_md(matrix: dict, cats) -> list[str]:
    lines = ["| human \\ judge | " + " | ".join(cats) + " |", "|---|" + "---|" * len(cats)]
    for r in cats:
        lines.append(f"| {r} | " + " | ".join(str(matrix[r][c]) for c in cats) + " |")
    return lines


def judge_kappa_markdown(purpose: str, reference: dict, results: dict, deciding_judge: str = "primary") -> str:
    lines = [f"# Judge vs human agreement -- {purpose} set", "",
             f"Generated: {datetime.now(timezone.utc).isoformat()}", "",
             f"Human reference: {reference['n_items']} items (status: {reference['status']}; "
             f"{reference.get('n_adjudicated', 0)} adjudicated).", ""]
    hh = reference.get("human_human")
    if hh:
        lines += [f"Human-human (pre-adjudication, n = {hh['n']}): 3-category kappa = {hh['kappa_3cat']}; "
                  f"binary kappa on both-non-ABSTAIN (n = {hh['n_both_non_abstain']}) = "
                  f"{hh['kappa_binary_non_abstain']}.", ""]
    lines += ["Acceptance (Amendment 4): unweighted 3-category Cohen's kappa (ABSTAIN / not hallucinated / "
              f"hallucinated), {deciding_judge} judge vs human reference, >= {ACCEPTANCE_KAPPA}. The binary kappa "
              "(hallucinated vs not, items both label non-ABSTAIN) is reported alongside and does not decide.", "",
              "| judge version | judge | n | 3-category kappa | acceptance | n both non-ABSTAIN | binary kappa | "
              "judge missing | judge unusable |", "|---|---|---|---|---|---|---|---|---|"]
    for (version, judge_id), r in results.items():
        k3 = r.get("kappa_3cat")
        accept = "-" if judge_id != deciding_judge else ("n/a" if k3 is None else "PASS" if k3 >= ACCEPTANCE_KAPPA else "below target")
        lines.append(f"| {version} | {judge_id} | {r['n']} | {k3} | {accept} | {r.get('n_both_non_abstain')} | "
                     f"{r.get('kappa_binary_non_abstain')} | {r['n_judge_missing']} | {r['n_judge_unusable']} |")
    for (version, judge_id), r in results.items():
        if not r["n"]:
            continue
        lines += ["", f"## {version} -- {judge_id} (n = {r['n']})", "", "3-category confusion matrix:", ""]
        lines += _matrix_md(r["confusion_3cat"], CATEGORIES)
        lines += ["", f"Binary confusion matrix, items both label non-ABSTAIN (n = {r['n_both_non_abstain']}):", ""]
        lines += _matrix_md(r["confusion_binary_non_abstain"], (CATEGORY_NOT_HALLUCINATED, CATEGORY_HALLUCINATED))
    return "\n".join(lines) + "\n"


def cmd_judge_kappa(args) -> int:
    from judge_agreement import load_judge_records

    purpose, pdir = args.purpose, purpose_dir(args.purpose)
    ref_path = pdir / "reference_adjudicated.json"
    if not ref_path.exists():
        raise AnnotationError(f"{ref_path} not found -- run `finalize-reference --purpose {purpose}` first")
    reference = json.loads(ref_path.read_text())
    key = load_verified_key(purpose, key_dir_from(args.key_dir))
    records = []
    for path in args.judge_output:  # one file at a time: a file never mixes judge versions, a set of files may
        records += load_judge_records([path])
    results = judge_vs_human(reference, key, records)
    out = Path(args.out) if args.out else pdir / "judge_vs_human_report.md"
    out.write_text(judge_kappa_markdown(purpose, reference, results))
    for (version, judge_id), r in results.items():
        log(f"{version} / {judge_id}: n={r['n']} kappa_3cat={r.get('kappa_3cat')} "
            f"binary(n={r.get('n_both_non_abstain')})={r.get('kappa_binary_non_abstain')}")
    log(f"Report -> {out}")
    return 0


# ---------------------------------------------------------------------------
# export-reannotation / intra-kappa (Amendment 11)
# ---------------------------------------------------------------------------

def select_reannotation(item_ids: list[str], purpose: str, seed: int = DEFAULT_SEED,
                        fraction: float = REANNOTATION_FRACTION) -> list[tuple[str, str]]:
    """[(new_id, original_id)] in the new sheet order: round(fraction * n) items drawn with a fixed seed,
    new ids assigned after a fresh shuffle (they carry no information about the original id or order)."""
    pool = sorted(item_ids)
    n = max(1, round(fraction * len(pool)))
    chosen = _rng(seed, purpose, "reannotation", "sample").sample(pool, n)
    _rng(seed, purpose, "reannotation", "order").shuffle(chosen)
    prefix = PURPOSES[purpose]["prefix"]
    width = max(2, len(str(n)))
    return [(f"{prefix}R-{i:0{width}d}", orig) for i, orig in enumerate(chosen, start=1)]


def days_since(iso: str, now: datetime | None = None) -> float:
    return ((now or datetime.now(timezone.utc)) - datetime.fromisoformat(iso)).total_seconds() / 86400


def cmd_export_reannotation(args, now: datetime | None = None, annotation_root: Path | None = None) -> int:
    purpose, annotator = args.purpose, args.annotator
    pdir, key_dir = purpose_dir(purpose, annotation_root), key_dir_from(args.key_dir)
    sheet = pdir / f"{annotator}.xlsx"
    comp_path = completion_path(purpose, annotator, annotation_root)
    if not comp_path.exists():
        raise AnnotationError(f"{comp_path.name} not found: the sheet was never validated as complete -- run "
                              f"`validate --xlsx {sheet}` first")
    completion = json.loads(comp_path.read_text())
    if completion["sheet_sha256"] != file_sha256(sheet):
        raise AnnotationError(f"{sheet.name} changed after it was recorded complete -- validate it again first")
    key = load_verified_key(purpose, key_dir, annotation_root)
    rows, problems = load_annotator(purpose, annotator, key, annotation_root)
    if problems:
        raise AnnotationError(f"{sheet.name} is not complete: {len(problems)} problem(s), e.g. {problems[0]}")
    waited = days_since(completion["completed_at_utc"], now)
    override = bool(args.override_min_days)
    if waited < REANNOTATION_MIN_DAYS and not override:
        raise AnnotationError(f"only {waited:.1f} days since {sheet.name} was completed "
                              f"({completion['completed_at_utc']}); re-annotation needs >= {REANNOTATION_MIN_DAYS} "
                              f"days. Override (recorded in the manifest): --override-min-days \"<reason>\"")
    out = pdir / f"reannotation_{annotator}.xlsx"
    if out.exists() and not args.force:
        raise AnnotationError(f"{out} already exists; pass --force to overwrite")

    pairs = select_reannotation(sorted(rows), purpose, args.seed)
    texts = _texts_from_sheet(sheet)
    write_annotation_xlsx(out, [{"item_id": new, **{c: texts[orig][c] or "" for c in TEXT_COLUMNS}}
                                for new, orig in pairs], annotator)
    mapping = {"purpose": purpose, "annotator": annotator, "pairs": [{"new_id": n, "item_id": o} for n, o in pairs]}
    key_dir.mkdir(parents=True, exist_ok=True)
    map_path = key_dir / f"{purpose}_{annotator}_reannotation_map_seed{args.seed}.json"
    map_path.write_text(json.dumps(mapping, indent=2) + "\n")
    manifest = {
        "purpose": purpose, "annotator": annotator, "seed": args.seed, "fraction": REANNOTATION_FRACTION,
        "n_items": len(pairs), "exported_at_utc": (now or datetime.now(timezone.utc)).isoformat(),
        "source_sheet": sheet.name, "source_sheet_sha256": completion["sheet_sha256"],
        "source_completed_at_utc": completion["completed_at_utc"], "days_since_completion": round(waited, 2),
        "min_days": REANNOTATION_MIN_DAYS, "override_min_days": override,
        "override_reason": args.override_min_days if override else None,
        "map_file_name": map_path.name, "map_sha256": file_sha256(map_path),
        "map_location_note": "kept OUTSIDE the repository (it links new ids to the first-pass items)",
    }
    (pdir / f"reannotation_{annotator}_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    log(f"Wrote {len(pairs)} items -> {out} ({waited:.1f} days after completion"
        f"{'; OVERRIDE: ' + args.override_min_days if override else ''})")
    log(f"Wrote the id map OUTSIDE the repo -> {map_path}")
    return 0


def intra_annotator(first: dict, second: dict, pairs: list[dict]) -> dict:
    """Agreement between the first pass (by original id) and the re-annotation (by new id)."""
    a = [first[p["item_id"]]["category"] for p in pairs]
    b = [second[p["new_id"]]["category"] for p in pairs]
    rep = agreement_report(a, b)
    rep["changes"] = [{"item_id": p["item_id"], "new_id": p["new_id"], "first": x, "second": y}
                      for p, x, y in zip(pairs, a, b) if x != y]
    return rep


def cmd_intra_kappa(args, annotation_root: Path | None = None) -> int:
    purpose, annotator = args.purpose, args.annotator
    pdir, key_dir = purpose_dir(purpose, annotation_root), key_dir_from(args.key_dir)
    man_path = pdir / f"reannotation_{annotator}_manifest.json"
    if not man_path.exists():
        raise AnnotationError(f"{man_path.name} not found -- run `export-reannotation` first")
    manifest = json.loads(man_path.read_text())
    map_path = key_dir / manifest["map_file_name"]
    if not map_path.exists() or file_sha256(map_path) != manifest["map_sha256"]:
        raise AnnotationError(f"re-annotation map {map_path} missing or not the one of this export")
    pairs = json.loads(map_path.read_text())["pairs"]
    key = load_verified_key(purpose, key_dir, annotation_root)
    first, p1 = load_annotator(purpose, annotator, key, annotation_root)
    second_rows, p2 = read_annotation_xlsx(pdir / f"reannotation_{annotator}.xlsx")
    second = {r["item_id"]: r for r in second_rows}
    expected = {k["item_id"]: k for k in key["items"]}
    for p in pairs:  # same texts as the first pass, every item labelled
        r = second.get(p["new_id"])
        if r is None:
            p2.append(f"re-annotation item {p['new_id']} is missing")
        elif r["text_sha256"] != expected[p["item_id"]]["text_sha256"]:
            p2.append(f"re-annotation item {p['new_id']}: text differs from the exported item")
    problems = p1 + p2
    if problems:
        log(f"[STOP] {len(problems)} problem(s):")
        for p in problems:
            log(f"  - {p}")
        return 1
    rep = intra_annotator(first, second, pairs)
    rep.update({"days_since_completion": manifest["days_since_completion"],
                "override_min_days": manifest["override_min_days"]})
    (pdir / f"intra_annotator_{annotator}.json").write_text(json.dumps(rep, indent=2) + "\n")
    log(f"Intra-annotator ({annotator}, n = {rep['n']}): 3-category kappa = {rep['kappa_3cat']}; "
        f"binary (n = {rep['n_both_non_abstain']}) = {rep['kappa_binary_non_abstain']}; "
        f"{len(rep['changes'])} label change(s)")
    return 0


# ---------------------------------------------------------------------------
# agreement (Amendment 11): every rater pair, bootstrap 95% CI, faithfulness, familiarity split
# ---------------------------------------------------------------------------

def judges_by_item(key: dict, judge_records: list[dict], judge_version: str) -> dict:
    """{judge_id: {item_id: record}} for one judge version, joined through the key by (run_id, question_id)."""
    by_rq = {(k["run_id"], k["question_id"]): k["item_id"] for k in key["items"]}
    out: dict[str, dict] = {}
    for r in judge_records:
        if r.get("prompt_version") != judge_version:
            continue
        item_id = by_rq.get((r.get("run_id"), r.get("question_id")))
        if item_id is not None:
            out.setdefault(r["judge_id"], {})[item_id] = r
    return out


def _paired(a: dict, b: dict, subset=None) -> tuple[list, list, int]:
    ids = sorted(i for i in set(a) | set(b) if subset is None or i in subset)
    both = [i for i in ids if a.get(i) is not None and b.get(i) is not None]
    return [a[i] for i in both], [b[i] for i in both], len(ids) - len(both)


def agreement_pairs(researcher: dict, judges: dict, annotator2: dict | None = None,
                    intra: dict | None = None) -> dict:
    """researcher / annotator2: {item_id: sheet row}; judges: {judge_id: {item_id: judge record}};
    intra: {"first": {item_id: category}, "second": {item_id: category}} (keyed by the original item id).
    Each pair: agreement_report on the items both raters label (others counted as n_excluded)."""
    cat = {"researcher": {i: r["category"] for i, r in researcher.items()}}
    for jid in ("primary", "secondary"):
        if jid in judges:
            cat[jid] = {i: category_of(r) for i, r in judges[jid].items()}
    if annotator2:
        cat["annotator2"] = {i: r["category"] for i, r in annotator2.items()}
    plan = [("primary_vs_researcher", "primary", "researcher"), ("secondary_vs_researcher", "secondary", "researcher"),
            ("primary_vs_secondary", "primary", "secondary"), ("annotator2_vs_researcher", "annotator2", "researcher")]
    pairs = {}
    for name, x, y in plan:
        if x in cat and y in cat:
            a, b, excluded = _paired(cat[x], cat[y])
            pairs[name] = {**agreement_report(a, b), "n_excluded": excluded}
    if intra:
        a, b, excluded = _paired(intra["first"], intra["second"])
        pairs["intra_researcher"] = {**agreement_report(a, b), "n_excluded": excluded}

    familiarity = {}
    for name, x in (("primary_vs_researcher", "primary"), ("secondary_vs_researcher", "secondary")):
        if x not in cat:
            continue
        for level in ("familiar", "not_familiar"):
            subset = {i for i, r in researcher.items() if r.get("domain_familiarity") == level}
            a, b, excluded = _paired(cat[x], cat["researcher"], subset)
            familiarity[f"{name}|{level}"] = {**agreement_report(a, b), "n_excluded": excluded}

    faithfulness = {}
    human_f = {i: r.get("faithfulness_level") for i, r in researcher.items() if r.get("has_context")}
    for jid in ("primary", "secondary"):
        if jid not in judges:
            continue
        from judge_prompt_v2_2 import faithfulness_of
        judge_f = {i: faithfulness_category(r.get("faithfulness", faithfulness_of(r))) for i, r in judges[jid].items()
                   if i in human_f}
        a, b, excluded = _paired(judge_f, human_f)
        faithfulness[f"{jid}_vs_researcher"] = {**kappa_report(a, b, FAITHFULNESS_LEVELS), "n_excluded": excluded}
    return {"pairs": pairs, "familiarity": familiarity, "faithfulness": faithfulness}


def _ci(ci) -> str:
    return f"[{ci[0]}, {ci[1]}]" if ci else "-"


def agreement_markdown(purpose: str, judge_version: str, rep: dict) -> str:
    lines = [f"# Agreement report -- {purpose} set, {judge_version}", "",
             f"Generated: {datetime.now(timezone.utc).isoformat()}. Unweighted Cohen's kappa; 95% CI = percentile "
             "bootstrap over item pairs (2,000 resamples, fixed seed). Acceptance (Amendment 4/11): primary judge vs "
             f"researcher, 3-category kappa >= {ACCEPTANCE_KAPPA}.", "",
             "| pair | n | excluded | 3-category kappa | 95% CI | n both non-ABSTAIN | binary kappa | 95% CI |",
             "|---|---|---|---|---|---|---|---|"]
    for section in ("pairs", "familiarity"):
        for name, r in rep[section].items():
            lines.append(f"| {name} | {r['n']} | {r['n_excluded']} | {r['kappa_3cat']} | {_ci(r['kappa_3cat_ci95'])} | "
                         f"{r['n_both_non_abstain']} | {r['kappa_binary_non_abstain']} | "
                         f"{_ci(r['kappa_binary_non_abstain_ci95'])} |")
    if rep["faithfulness"]:
        lines += ["", "## Faithfulness (judge ratio -> all / partly / not supported, vs researcher)", "",
                  "| pair | n | excluded | kappa | 95% CI | observed agreement |", "|---|---|---|---|---|---|"]
        for name, r in rep["faithfulness"].items():
            lines.append(f"| {name} | {r['n']} | {r['n_excluded']} | {r['kappa']} | {_ci(r['kappa_ci95'])} | "
                         f"{r['observed_agreement']} |")
    for name, r in rep["pairs"].items():
        if r["n"]:
            lines += ["", f"## {name} (n = {r['n']})", "", "3-category confusion matrix (rows: first rater):", ""]
            lines += _matrix_md(r["confusion_3cat"], CATEGORIES)
    return "\n".join(lines) + "\n"


def cmd_agreement(args, annotation_root: Path | None = None) -> int:
    from judge_agreement import load_judge_records

    purpose, pdir = args.purpose, purpose_dir(args.purpose, annotation_root)
    key = load_verified_key(purpose, key_dir_from(args.key_dir), annotation_root)
    researcher, problems = load_annotator(purpose, "annotator1", key, annotation_root)
    annotator2 = None
    if sheet_started(pdir / "annotator2.xlsx"):
        annotator2, p2 = load_annotator(purpose, "annotator2", key, annotation_root)
        problems += p2
    if problems:
        log(f"[STOP] {len(problems)} problem(s) in the annotation files -- fix them first:")
        for p in problems[:20]:
            log(f"  - {p}")
        return 1
    records = []
    for path in args.judge_output:
        records += load_judge_records([path])
    judges = judges_by_item(key, records, args.judge_version)
    rep = agreement_pairs(researcher, judges, annotator2)
    intra_path = pdir / "intra_annotator_annotator1.json"  # written by `intra-kappa`
    if intra_path.exists():
        rep["pairs"]["intra_researcher"] = {**json.loads(intra_path.read_text()), "n_excluded": 0}
    out = Path(args.out) if args.out else pdir / f"agreement_{args.judge_version}.md"
    out.write_text(agreement_markdown(purpose, args.judge_version, rep))
    out.with_suffix(".json").write_text(json.dumps(rep, indent=2, default=str) + "\n")
    for name, r in rep["pairs"].items():
        log(f"{name}: n={r['n']} kappa_3cat={r['kappa_3cat']} {_ci(r.get('kappa_3cat_ci95'))}")
    log(f"Report -> {out}")
    return 0


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)

    def common(p):
        p.add_argument("--purpose", choices=sorted(PURPOSES), required=True)
        p.add_argument("--key-dir", default=None,
                       help=f"Directory of the key file, outside the repo (default: $GRAPHRAG_ANNOTATION_KEY_DIR "
                            f"or {DEFAULT_KEY_DIR})")

    p = sub.add_parser("export")
    common(p)
    p.add_argument("--run-files", nargs="*", default=[],
                   help="Run JSONL files to sample from (calibration: defaults to the calibration registry)")
    p.add_argument("--seed", type=int, default=DEFAULT_SEED)
    p.add_argument("--n-total", type=int, default=None,
                   help="only for a purpose without a fixed per-stratum count (calibration: 20, confirmation: 7)")
    p.add_argument("--questions-parquet", default=None)
    p.add_argument("--force", action="store_true", help="Overwrite an existing export (invalidates annotation in progress)")
    p.set_defaults(func=cmd_export)

    p = sub.add_parser("validate")
    common(p)
    p.add_argument("--xlsx", required=True)
    p.add_argument("--annotator", choices=["annotator1", "annotator2"], default=None)
    p.set_defaults(func=cmd_validate)

    p = sub.add_parser("human-kappa")
    common(p)
    p.add_argument("--force", action="store_true", help="Rewrite adjudication.xlsx even if it exists")
    p.set_defaults(func=cmd_human_kappa)

    p = sub.add_parser("finalize-reference")
    common(p)
    p.set_defaults(func=cmd_finalize_reference)

    p = sub.add_parser("judge-kappa")
    common(p)
    p.add_argument("--judge-output", nargs="+", required=True)
    p.add_argument("--out", default=None)
    p.set_defaults(func=cmd_judge_kappa)

    p = sub.add_parser("export-reannotation")
    common(p)
    p.add_argument("--annotator", choices=["annotator1", "annotator2"], default="annotator1")
    p.add_argument("--seed", type=int, default=DEFAULT_SEED)
    p.add_argument("--override-min-days", default=None, metavar="REASON",
                   help=f"Export before {REANNOTATION_MIN_DAYS} days have passed; the reason is recorded in the manifest")
    p.add_argument("--force", action="store_true")
    p.set_defaults(func=cmd_export_reannotation)

    p = sub.add_parser("intra-kappa")
    common(p)
    p.add_argument("--annotator", choices=["annotator1", "annotator2"], default="annotator1")
    p.set_defaults(func=cmd_intra_kappa)

    p = sub.add_parser("agreement")
    common(p)
    p.add_argument("--judge-output", nargs="+", required=True)
    p.add_argument("--judge-version", required=True, help="e.g. judge-v2.2")
    p.add_argument("--out", default=None)
    p.set_defaults(func=cmd_agreement)

    args = parser.parse_args(argv)
    try:
        return args.func(args)
    except AnnotationError as e:
        log(f"[ERROR] {e}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
