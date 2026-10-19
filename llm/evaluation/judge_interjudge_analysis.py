"""
judge_interjudge_analysis.py
----------------------------
Inter-judge agreement and hard-case identification for the panel judges
(llama, deepseek, qwen) on the calibration set. Descriptive only: kappa and
percentages, no inferential tests. Reads the SAVED panel outputs; never calls
a judge API; never writes a human label file.

Per judge version (judge-v2.3 = main result, judge-v2.2 = comparison), each
against the frozen human labels (annotator1.xlsx, guideline v1):

1. Inter-judge agreement (judge_binary.category_of, the same 3-category label
   as judge_panel_agreement.py):
     - binary "hallucinated" (PRIMARY): ABSTAIN excluded (an item or pair is
       left out when any rater in it says ABSTAIN); companion row with ABSTAIN
       counted as "not hallucinated"; 3-category rows as a supplement;
     - pairwise unweighted Cohen's kappa + raw percent agreement per pair;
     - Fleiss' kappa for the whole panel, on COMPLETE items only (all three
       judges usable), with n;
     - prevalence of "hallucinated" per judge (and the human, for reference);
     - severity: a DISTRIBUTION table per judge vs the human
       (FAKTUAL / SEBAGIAN / PENUH / UNSPECIFIED / ABSTAIN / missing), no kappa.
   Missing / failed judge records are listed per judge, never imputed.

2. Item categories (3-category labels, ABSTAIN / NOT_HALLUCINATED / HALLUCINATED):
     A  panel unanimous and equal to the human      -> easy
     B  panel unanimous but different from the human -> re-check the human label
     C  panel split                                  -> ambiguous / hard
     D  a judge has no usable label (failed / parse error / missing)

3. Review sheet (xlsx) for B and C items, blind to the condition: the item text
   and the human label first, empty reviewer columns, then each judge's label
   and reasoning in the right-most column groups.

4. One summary markdown for all versions; the hand-written category-C pattern
   notes (--patterns-md) are included verbatim.

CLI (from llm/evaluation; joins through the verified key outside the repo):
    python3 judge_interjudge_analysis.py --purpose calibration \
        --versions judge-v2.3 judge-v2.2 --patterns-md results/interjudge_patterns_C_calibration.md
"""

import argparse
import json
import sys
from collections import Counter
from datetime import datetime
from itertools import combinations
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from judge_binary import (CATEGORIES, CATEGORY_ABSTAIN, CATEGORY_HALLUCINATED, CATEGORY_NOT_HALLUCINATED,  # noqa: E402
                          ERROR_VERDICTS, category_of, cohen_kappa, is_unusable, severity_of)
from judge_panel_agreement import SENSITIVITY_LABELS  # noqa: E402

JUDGES = ("llama", "deepseek", "qwen")
BINARY = (CATEGORY_NOT_HALLUCINATED, CATEGORY_HALLUCINATED)
MODES = ("abstain_excluded", "abstain_counted", "three_category")
MODE_TITLE = {"abstain_excluded": "binary, ABSTAIN excluded (PRIMARY)",
              "abstain_counted": "binary, ABSTAIN counted as not hallucinated",
              "three_category": "3-category (ABSTAIN / NOT_HALLUCINATED / HALLUCINATED)"}
SEVERITY_COLUMNS = ("FAKTUAL", "SEBAGIAN", "PENUH", "UNSPECIFIED", "ABSTAIN", "missing")
CATEGORY_DESC = {"A": "panel unanimous, equal to the human (easy)",
                 "B": "panel unanimous, different from the human (re-check the human label)",
                 "C": "panel split (ambiguous / hard)",
                 "D": "incomplete: a judge has no usable label"}
XLSX_CELL_MAX = 32000  # Excel's limit is 32,767 characters per cell
RESULTS = Path(__file__).resolve().parent / "results"


# ---------------------------------------------------------------- statistics

def to_mode(label: str | None, mode: str) -> str | None:
    """A 3-category label on the scale of `mode`; None = not usable in this mode."""
    if label is None:
        return None
    if mode == "abstain_excluded":
        return None if label == CATEGORY_ABSTAIN else label
    if mode == "abstain_counted":
        return CATEGORY_NOT_HALLUCINATED if label == CATEGORY_ABSTAIN else label
    return label


def mode_categories(mode: str) -> tuple:
    return CATEGORIES if mode == "three_category" else BINARY


def fleiss_kappa(rows: list[list[str]], categories) -> float | None:
    """Fleiss' kappa; every row holds the labels of the same number (>= 2) of raters for one item. None when
    undefined (no items, or expected agreement == 1: every rating in one category)."""
    if not rows:
        return None
    m = len(rows[0])
    if m < 2 or any(len(r) != m for r in rows):
        raise ValueError("every item needs the same number (>= 2) of ratings")
    n_items = len(rows)
    counts = [Counter(r) for r in rows]
    p_bar = sum((sum(c[k] ** 2 for k in categories) - m) / (m * (m - 1)) for c in counts) / n_items
    p_e = sum((sum(c[k] for c in counts) / (n_items * m)) ** 2 for k in categories)
    if p_e >= 1.0:
        return None
    return (p_bar - p_e) / (1 - p_e)


def _rate(x: int, n: int) -> dict:
    return {"x": x, "n": n, "rate": round(x / n, 4) if n else None}


def _round(k: float | None) -> float | None:
    return round(k, 4) if k is not None else None


def pairwise(raters: dict, items: list[str], mode: str) -> dict:
    """{"a_vs_b": {n, agree, kappa}} over the items both raters have a label for in `mode`."""
    out = {}
    for a, b in combinations(JUDGES, 2):
        pairs = [(to_mode(raters[a][i], mode), to_mode(raters[b][i], mode)) for i in items]
        pairs = [(x, y) for x, y in pairs if x is not None and y is not None]
        la, lb = [p[0] for p in pairs], [p[1] for p in pairs]
        out[f"{a}_vs_{b}"] = {"n": len(pairs), "agree": _rate(sum(x == y for x, y in pairs), len(pairs)),
                              "kappa": _round(cohen_kappa(la, lb, mode_categories(mode)))}
    return out


def fleiss(raters: dict, items: list[str], mode: str) -> dict:
    """Fleiss' kappa over the items where every judge has a label in `mode` (complete items only)."""
    rows = []
    for i in items:
        labels = [to_mode(raters[j][i], mode) for j in JUDGES]
        if all(lab is not None for lab in labels):
            rows.append(labels)
    unanimous = sum(len(set(r)) == 1 for r in rows)
    return {"n": len(rows), "kappa": _round(fleiss_kappa(rows, mode_categories(mode))),
            "unanimous": _rate(unanimous, len(rows))}


def prevalence(labels: dict, items: list[str]) -> dict:
    """Share labelled HALLUCINATED, per mode (denominator: the rater's usable labels in that mode)."""
    out = {}
    for mode in ("abstain_excluded", "abstain_counted"):
        vals = [to_mode(labels[i], mode) for i in items]
        vals = [v for v in vals if v is not None]
        out[mode] = _rate(vals.count(CATEGORY_HALLUCINATED), len(vals))
    vals = [labels[i] for i in items if labels[i] is not None]
    out["abstain"] = _rate(vals.count(CATEGORY_ABSTAIN), len(vals))
    return out


def severity_bucket(category: str | None, severity: str | None) -> str:
    if category is None:
        return "missing"
    if category == CATEGORY_ABSTAIN:
        return "ABSTAIN"
    if category == CATEGORY_NOT_HALLUCINATED:
        return "FAKTUAL"
    return severity if severity in ("SEBAGIAN", "PENUH") else "UNSPECIFIED"


def categorize(judge_labels: list[str | None], human: str | None) -> str:
    """A / B / C / D on the 3-category scale (see the module docstring)."""
    if any(lab is None for lab in judge_labels) or human is None:
        return "D"
    if len(set(judge_labels)) > 1:
        return "C"
    return "A" if judge_labels[0] == human else "B"


# ---------------------------------------------------------------- loading

def judge_view(key_items: list[dict], records: list[dict], failures: list[dict]) -> tuple[dict, dict, dict]:
    """(labels, records, missing): labels {judge: {item_id: category | None}}, records {judge: {item_id: record}},
    missing {judge: {item_id: reason}} -- every item without a usable label, with the reason."""
    from judge_panel_config import PANEL_ID, PANEL_JUDGES

    by_rq = {(k["run_id"], int(k["question_id"])): k["item_id"] for k in key_items}
    model_to_id = {j.model: j.judge_id for j in PANEL_JUDGES}
    recs = {j: {} for j in JUDGES}
    unmatched = []
    for rec in records:
        item = by_rq.get((rec.get("run_id"), int(rec.get("question_id", -1))))
        jid = model_to_id.get(rec.get("judge_model"))
        if item is None or jid not in recs or rec.get("panel_id") != PANEL_ID:
            unmatched.append(rec.get("item_id"))
            continue
        if item in recs[jid]:
            raise ValueError(f"duplicate record for {jid} / {item}")
        recs[jid][item] = rec
    if unmatched:
        raise ValueError(f"{len(unmatched)} judge record(s) do not match the key / panel: {unmatched[:5]}")
    failed = {}
    for rec in failures:
        item = by_rq.get((rec.get("run_id"), int(rec.get("question_id", -1))))
        failed[(model_to_id.get(rec.get("judge_model")), item)] = rec.get("error") or "failed"
    labels, missing = {j: {} for j in JUDGES}, {j: {} for j in JUDGES}
    for k in key_items:
        i = k["item_id"]
        for j in JUDGES:
            rec = recs[j].get(i)
            labels[j][i] = category_of(rec) if rec is not None else None
            if rec is None:
                missing[j][i] = f"no record; failures file: {failed.get((j, i), 'none')}"
            elif is_unusable(rec):
                missing[j][i] = f"unusable record ({rec.get('error') or 'parse/call failure'})"
    return labels, recs, missing


def load_human_text(sheet: Path) -> dict:
    """{item_id: {question, accepted_answer, candidate_answer, context, wrong_claims, notes}} read-only."""
    from openpyxl import load_workbook

    ws = load_workbook(sheet, read_only=True, data_only=True)["anotasi"]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    want = ("question", "accepted_answer", "candidate_answer", "context", "wrong_claims", "notes")
    out = {}
    for r in rows:
        d = dict(zip(header, r))
        if d.get("item_id"):
            out[str(d["item_id"]).strip()] = {c: "" if d.get(c) is None else str(d[c]) for c in want}
    return out


# ---------------------------------------------------------------- analysis

def analyze(version: str, key_items: list[dict], human: dict, human_sev: dict, records: list[dict],
            failures: list[dict]) -> dict:
    items = sorted(k["item_id"] for k in key_items)
    labels, recs, missing = judge_view(key_items, records, failures)
    judge_sev = {j: {i: severity_bucket(labels[j][i], severity_of(recs[j][i]) if i in recs[j] else None)
                     for i in items} for j in JUDGES}
    human_bucket = {i: severity_bucket(human[i], human_sev.get(i)) for i in items}

    cats = {i: categorize([labels[j][i] for j in JUDGES], human[i]) for i in items}
    sens = SENSITIVITY_LABELS.get("calibration", {})
    sens_human = {**human, **{i: v for i, v in sens.items() if i in human}}
    cats_sens = {i: categorize([labels[j][i] for j in JUDGES], sens_human[i]) for i in items}

    return {
        "prompt_version": version,
        "n_items": len(items),
        "missing": {j: dict(sorted(missing[j].items())) for j in JUDGES},
        "pairwise": {m: pairwise(labels, items, m) for m in MODES},
        "fleiss": {m: fleiss(labels, items, m) for m in MODES},
        "prevalence": {**{j: prevalence(labels[j], items) for j in JUDGES}, "human": prevalence(human, items)},
        "severity_distribution": {r: dict(Counter(v.values())) for r, v in
                                  {**judge_sev, "human": human_bucket}.items()},
        "judge_sebagian_count": sum(list(judge_sev[j].values()).count("SEBAGIAN") for j in JUDGES),
        "categories": {c: sorted(i for i in items if cats[i] == c) for c in "ABCD"},
        "categories_sensitivity": {c: sum(1 for i in items if cats_sens[i] == c) for c in "ABCD"},
        "sensitivity_items": sorted(sens),
        "per_item": {i: {"category": cats[i], "human": human[i], "human_severity": human_sev.get(i),
                         **{j: labels[j][i] for j in JUDGES}} for i in items},
        "_records": recs,  # not serialised
    }


# ---------------------------------------------------------------- outputs

def _cell(text: str) -> str:
    text = "" if text is None else str(text)
    return text if len(text) <= XLSX_CELL_MAX else text[:XLSX_CELL_MAX] + "\n[... truncated for Excel]"


def _claims_text(rec: dict | None, errors: bool) -> str:
    if rec is None:
        return ""
    lines = []
    for c in rec.get("claims") or []:
        verdict = str(c.get("verdict", "")).upper()
        if (verdict in ERROR_VERDICTS) != errors:
            continue
        sev = f" [{c.get('severity')}]" if c.get("severity") else ""
        lines.append(f"- {verdict}{sev}: {c.get('claim', '')}\n  evidence: {c.get('evidence', '')}")
    return "\n".join(lines)


def _human_label_text(category: str | None, severity: str | None) -> str:
    if category == CATEGORY_HALLUCINATED:
        return f"HALLUCINATED ({severity})"
    return category or "missing"


def write_review_sheet(path: Path, rep: dict, texts: dict) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    base = [("item_id", 10), ("kategori", 9), ("question", 60), ("jawaban_LLM (candidate_answer)", 60),
            ("referensi_SORD (accepted_answer)", 60), ("context [K1..K5]", 45),
            ("label_manusia_v1", 18), ("wrong_claims_manusia_v1", 40), ("notes_manusia_v1", 30),
            ("post_freeze_edit", 10)]
    review = [("keputusan_review", 20), ("aturan_guideline", 20), ("catatan", 35)]
    per_judge = [("kategori", 18), ("label", 18), ("severity", 11), ("reasoning", 55), ("klaim_error", 55),
                 ("klaim_lain", 55)]
    groups = [("Item (nilai dulu sebelum melihat judge)", base), ("Keputusan reviewer (isi manual)", review)] + \
             [(f"Judge: {j}", per_judge) for j in JUDGES]

    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical="top")
    fills = ["DDEBF7", "FFF2CC", "E2EFDA", "FCE4D6", "EDEDED"]
    col = 1
    for g, (title, cols) in enumerate(groups):
        ws.cell(row=1, column=col, value=title).font = bold
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(cols) - 1)
        for name, width in cols:
            c = ws.cell(row=2, column=col, value=name)
            c.font, c.alignment = bold, wrap
            c.fill = PatternFill("solid", fgColor=fills[g])
            ws.column_dimensions[get_column_letter(col)].width = width
            col += 1
        ws.cell(row=1, column=col - len(cols)).fill = PatternFill("solid", fgColor=fills[g])

    sens = set(rep["sensitivity_items"])
    review_items = [(c, i) for c in "BC" for i in rep["categories"][c]]
    for r, (cat, i) in enumerate(review_items, start=3):
        t, info = texts[i], rep["per_item"][i]
        values = [i, cat, t["question"], t["candidate_answer"], t["accepted_answer"], t["context"],
                  _human_label_text(info["human"], info["human_severity"]), t["wrong_claims"], t["notes"],
                  "Y" if i in sens else "N", "", "", ""]
        for j in JUDGES:
            rec = rep["_records"][j].get(i)
            values += [info[j] or "missing", (rec or {}).get("label", ""),
                       severity_of(rec) or "" if rec else "", (rec or {}).get("reasoning", ""),
                       _claims_text(rec, True), _claims_text(rec, False)]
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=_cell(v)).alignment = wrap
    ws.freeze_panes = "C3"

    notes = wb.create_sheet("petunjuk")
    for line in (f"Review sheet {rep['prompt_version']} (panel-v1) -- calibration set, categories B and C.",
                 "B = panel unanimous but different from the human label; C = panel split. 3-category labels.",
                 "Condition / run is hidden (blind-v3). Judge columns are on the right: fill the reviewer "
                 "columns before reading them.",
                 "label_manusia_v1 = frozen annotator1.xlsx (guideline v1). post_freeze_edit = Y for the "
                 "items with a post-freeze edit (docs/ANNOTATION_POSTFREEZE_RECONSIDERATIONS.md).",
                 "This sheet is a working copy: human label files are never changed by this script."):
        notes.append([line])
    notes.column_dimensions["A"].width = 120
    wb.properties.created = datetime(2026, 1, 1)
    wb.save(path)
    _normalize_xlsx(path)
    return len(review_items)


def _normalize_xlsx(path: Path) -> None:
    """Byte-reproducible xlsx: openpyxl stamps the save time into docProps/core.xml and every zip entry."""
    import re
    import zipfile

    fixed = "2026-01-01T00:00:00Z"
    with zipfile.ZipFile(path) as z:
        entries = [(info.filename, z.read(info.filename)) for info in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            if name == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)",
                              rb"\g<1>" + fixed.encode() + rb"\g<2>", data)
            z.writestr(zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0)), data, zipfile.ZIP_DEFLATED)


def _fmt_rate(r: dict) -> str:
    return f"{r['x']}/{r['n']} ({100 * r['rate']:.1f}%)" if r["n"] else "0/0 (n/a)"


def _fmt_k(k) -> str:
    return "undefined" if k is None else f"{k:.3f}"


def version_markdown(rep: dict) -> str:
    v = rep["prompt_version"]
    out = [f"## {v}", ""]
    for mode in MODES:
        out += [f"### Inter-judge kappa -- {MODE_TITLE[mode]}", "",
                "| pair | n | raw agreement | Cohen's kappa |", "|---|---|---|---|"]
        for pair, d in rep["pairwise"][mode].items():
            out.append(f"| {pair.replace('_vs_', ' vs ')} | {d['n']} | {_fmt_rate(d['agree'])} | "
                       f"{_fmt_k(d['kappa'])} |")
        f = rep["fleiss"][mode]
        out += [f"| **Fleiss' kappa (panel, complete items)** | {f['n']} | unanimous {_fmt_rate(f['unanimous'])} "
                f"| **{_fmt_k(f['kappa'])}** |", ""]
        if mode == "abstain_excluded":
            out += ["Matrix (Cohen's kappa; raw agreement):", "",
                    "| | " + " | ".join(JUDGES) + " |", "|---|" + "---|" * len(JUDGES)]
            for a in JUDGES:
                cells = []
                for b in JUDGES:
                    if a == b:
                        cells.append("--")
                        continue
                    d = rep["pairwise"][mode].get(f"{a}_vs_{b}") or rep["pairwise"][mode][f"{b}_vs_{a}"]
                    cells.append(f"{_fmt_k(d['kappa'])} ({100 * d['agree']['rate']:.1f}%, n={d['n']})")
                out.append(f"| {a} | " + " | ".join(cells) + " |")
            out.append("")
    out += ["### Prevalence of \"hallucinated\" (and ABSTAIN) per rater", "",
            "| rater | H, ABSTAIN excluded | H, ABSTAIN counted | ABSTAIN |", "|---|---|---|---|"]
    for r, p in rep["prevalence"].items():
        out.append(f"| {r} | {_fmt_rate(p['abstain_excluded'])} | {_fmt_rate(p['abstain_counted'])} | "
                   f"{_fmt_rate(p['abstain'])} |")
    never_h = [j for j in JUDGES if rep["prevalence"][j]["abstain_counted"]["x"] == 0]
    if never_h:
        out += ["", f"{', '.join(never_h)} labelled no item HALLUCINATED: every binary Cohen's kappa involving "
                    f"{'it' if len(never_h) == 1 else 'them'} is 0 by construction (expected agreement equals "
                    "observed agreement on the other rater's NOT_HALLUCINATED items), even where raw agreement is "
                    "high. The binary kappa values reflect the low prevalence of HALLUCINATED, not only "
                    "disagreement."]
    out += ["", "### Severity distribution (no kappa)", "",
            "| rater | " + " | ".join(SEVERITY_COLUMNS) + " |", "|---|" + "---|" * len(SEVERITY_COLUMNS)]
    for r, d in rep["severity_distribution"].items():
        out.append(f"| {r} | " + " | ".join(str(d.get(c, 0)) for c in SEVERITY_COLUMNS) + " |")
    out += ["", f"Judge SEBAGIAN labels in this run: {rep['judge_sebagian_count']}. Severity of a judge's "
                "HALLUCINATED label = PENUH if any error claim is CORE, SEBAGIAN if every error claim is MINOR "
                "(judge_binary.severity_of).", "",
            "### Missing / unusable judge labels (not imputed)", ""]
    any_missing = False
    for j in JUDGES:
        for i, reason in rep["missing"][j].items():
            out.append(f"- {j} / {i}: {reason}")
            any_missing = True
    if not any_missing:
        out.append("- none: every item has a usable label from all three judges")
    out += ["", "### Item categories (3-category labels vs frozen annotator1.xlsx)", "",
            "| category | meaning | n | items |", "|---|---|---|---|"]
    for c in "ABCD":
        ids = rep["categories"][c]
        shown = ", ".join(ids) if c != "A" else "(see JSON)"
        out.append(f"| {c} | {CATEGORY_DESC[c]} | {len(ids)} ({100 * len(ids) / rep['n_items']:.1f}%) | "
                   f"{shown} |")
    s = rep["categories_sensitivity"]
    out += ["", f"SENSITIVITY (human labels with the post-freeze edits {', '.join(rep['sensitivity_items'])}; "
                f"not the frozen labels): A={s['A']}, B={s['B']}, C={s['C']}, D={s['D']}.", ""]
    if rep["categories"]["C"]:
        out += ["Category C items (3-category labels):", "",
                "| item | human | " + " | ".join(JUDGES) + " |", "|---|---|" + "---|" * len(JUDGES)]
        for i in rep["categories"]["C"]:
            p = rep["per_item"][i]
            out.append(f"| {i} | {_human_label_text(p['human'], p['human_severity'])} | "
                       + " | ".join(p[j] for j in JUDGES) + " |")
        out.append("")
    return "\n".join(out)


def summary_markdown(reports: list[dict], patterns: str | None, sources: dict) -> str:
    main_v = reports[0]["prompt_version"]
    out = ["# Inter-judge agreement and hard cases -- calibration set", "",
           f"Panel panel-v1 (llama = Llama-3.3-70B-Instruct-Turbo, deepseek = DeepSeek-V3.2, qwen = "
           f"Qwen2.5-72B-Instruct). Main result: {main_v}; comparison: "
           f"{', '.join(r['prompt_version'] for r in reports[1:]) or '-'}. Human reference: frozen "
           "annotator1.xlsx (guideline v1), 60 items. Descriptive statistics only (kappa, percentages); no "
           "inferential tests. Saved judge outputs only; no judge API call.", "",
           "Conventions:", "",
           "- Labels: judge_binary.category_of (derived from the judge's claims), the same 3-category label as "
           "judge_panel_agreement.py.",
           "- PRIMARY = binary \"hallucinated\" with ABSTAIN excluded (an item/pair is left out when a rater in "
           "it says ABSTAIN); companion = ABSTAIN counted as not hallucinated; 3-category as a supplement.",
           "- Cohen's kappa is unweighted, pairwise, over the items both judges labelled; Fleiss' kappa only on "
           "complete items (all three judges usable in that mode), n given.",
           "- Kappa can be low while raw agreement is high when one label is rare (see prevalence); \"undefined\" "
           "= both raters used one single category.",
           "- Severity: distribution only, no inter-judge kappa.",
           "- Categories A-D use the 3-category labels; failed / unparsed judge records are category D and are "
           "never imputed.", "",
           "## Overview", "",
           "| version | Fleiss' kappa binary excl. (n) | Fleiss' kappa binary counted (n) | Fleiss' kappa 3-cat (n) "
           "| A | B | C | D | judge SEBAGIAN |", "|---|---|---|---|---|---|---|---|---|"]
    for r in reports:
        f, c = r["fleiss"], r["categories"]
        out.append(f"| {r['prompt_version']} | {_fmt_k(f['abstain_excluded']['kappa'])} ({f['abstain_excluded']['n']}) "
                   f"| {_fmt_k(f['abstain_counted']['kappa'])} ({f['abstain_counted']['n']}) "
                   f"| {_fmt_k(f['three_category']['kappa'])} ({f['three_category']['n']}) "
                   f"| {len(c['A'])} | {len(c['B'])} | {len(c['C'])} | {len(c['D'])} | {r['judge_sebagian_count']} |")
    n_sebagian = sum(r["judge_sebagian_count"] for r in reports)
    out += ["", f"Judge SEBAGIAN labels across these runs: {n_sebagian}. "
                + ("No judge produced a SEBAGIAN label, so judge severity is PENUH-or-nothing while the human "
                   "labels hold both SEBAGIAN and PENUH; severity is therefore shown as a distribution only. "
                   if n_sebagian == 0 else "")
                + "The judge prompt was not changed for this analysis.", ""]
    for r in reports:
        out += [version_markdown(r), ""]
    out += ["## Patterns in category C", ""]
    out += [patterns.strip() if patterns else "_(no pattern notes given: pass --patterns-md)_", ""]
    out += ["## Sources", ""] + [f"- {k}: `{v}`" for k, v in sources.items()] + [""]
    return "\n".join(out)


# ---------------------------------------------------------------- CLI

def main(argv=None) -> int:
    from annotation_tool import (AnnotationError, assert_annotation_frozen, key_dir_from, load_annotator,
                                 load_verified_key, purpose_dir)
    from judge_panel import failures_path_of, read_jsonl

    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--purpose", choices=["calibration"], default="calibration")
    p.add_argument("--versions", nargs="+", default=["judge-v2.3", "judge-v2.2"],
                   help="judge prompt versions; the first one is the main result")
    p.add_argument("--patterns-md", default=None, help="hand-written category-C pattern notes (included verbatim)")
    p.add_argument("--key-dir", default=None)
    args = p.parse_args(argv)

    try:
        frozen = assert_annotation_frozen(args.purpose, "annotator1")
    except AnnotationError as e:
        print(f"[ERROR] {e}")
        return 1
    print(f"[freeze] {frozen['file_name']} matches its freeze record (sha256 {frozen['sha256'][:12]}...)")
    key = load_verified_key(args.purpose, key_dir_from(args.key_dir))
    rows, problems = load_annotator(args.purpose, "annotator1", key)
    if problems:
        print(f"[STOP] {len(problems)} problem(s) in the annotation sheet")
        return 1
    sheet = purpose_dir(args.purpose) / "annotator1.xlsx"
    human = {k["item_id"]: rows[k["item_id"]]["category"] for k in key["items"]}
    human_sev = {k["item_id"]: rows[k["item_id"]]["severity"] for k in key["items"]}
    texts = load_human_text(sheet)

    reports, sources = [], {"human labels": str(sheet.relative_to(Path.cwd()) if sheet.is_relative_to(Path.cwd())
                                                 else sheet)}
    for v in args.versions:
        judge_path = RESULTS / f"panel_{args.purpose}_{v}.jsonl"
        records, failures = read_jsonl(judge_path), read_jsonl(failures_path_of(judge_path))
        if not records:
            print(f"[ERROR] no records in {judge_path}")
            return 1
        found = {r.get("prompt_version") for r in records + failures}
        if found != {v}:
            print(f"[ERROR] {judge_path.name} holds prompt versions {sorted(map(str, found))}, expected {v}")
            return 1
        rep = analyze(v, key["items"], human, human_sev, records, failures)
        stem = f"interjudge_{args.purpose}_{v}"
        (RESULTS / f"{stem}.json").write_text(
            json.dumps({k: val for k, val in rep.items() if not k.startswith("_")}, indent=2) + "\n")
        (RESULTS / f"{stem}.md").write_text(f"# Inter-judge agreement -- {args.purpose}\n\n" + version_markdown(rep))
        n_review = write_review_sheet(RESULTS / f"review_hard_cases_{v}.xlsx", rep, texts)
        sources[v] = f"results/{judge_path.name}" + (f" (+ {failures_path_of(judge_path).name})" if failures else "")
        f = rep["fleiss"]["abstain_excluded"]
        print(f"{v}: Fleiss binary excl={f['kappa']} (n={f['n']}) | categories "
              + " ".join(f"{c}={len(rep['categories'][c])}" for c in "ABCD") + f" | review rows={n_review}")
        reports.append(rep)

    patterns = Path(args.patterns_md).read_text() if args.patterns_md else None
    out = RESULTS / f"interjudge_summary_{args.purpose}.md"
    out.write_text(summary_markdown(reports, patterns, sources))
    print(f"Summary -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
