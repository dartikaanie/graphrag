"""Amendment 11 annotation sheet: context column, blind-v3 candidate, evidence-backed wrong claims,
faithfulness / approach / familiarity / uncertain columns, annotator_id per row."""

import json

import pytest

import annotation_tool as at

CTX = [{"question_id": 7, "answer_id": 70, "chunk_text": "Q: dates\nA: use fromisoformat"}]


def _items():
    items = []
    for label, cond in (("A [dev]", "A"), ("B-plain [dev]", "B"), ("B-grounded [dev]", "B")):
        for q in range(1000, 1050):
            items.append({"run_id": f"{cond}-{label[:6]}", "run_label": label, "condition": cond,
                          "source_file": f"/x/{label}.jsonl", "config_hash": "h", "question_id": q,
                          "title": f"Title {q}", "ground_truth_answer_html": f"<p>Reference {q}</p>",
                          "candidate": (f"Answer {q} [SO-7]. The retrieved sources do not cover tz." if cond != "A"
                                        else f"Answer {q}."),
                          "retrieved_context": CTX if cond != "A" else [], "has_context": cond != "A"})
    return items


@pytest.fixture
def sheet(tmp_path, monkeypatch):
    monkeypatch.setattr(at, "ANNOTATION_ROOT", tmp_path / "annotation")
    export = at.build_export(_items(), {q: f"<p>Body {q}</p>" for q in range(1000, 1050)}, "calibration", 42)
    pdir = at.purpose_dir("calibration")
    at.write_annotation_xlsx(pdir / "annotator1.xlsx", export["annotator1"], "annotator1")
    at.write_annotation_xlsx(pdir / "annotator2.xlsx", export["annotator2"], "annotator2")
    return export, pdir


def test_export_shows_context_only_for_retrieval_answers_and_blind_v3(sheet):
    export, _ = sheet
    by_cond = {}
    for k, row in zip(export["key_items"], export["annotator1"]):
        assert k["item_id"] == row["item_id"]
        by_cond.setdefault(k["condition"], []).append((k, row))
    for k, row in by_cond["A"]:
        assert row["context"] == "" and k["has_context"] is False
    for k, row in by_cond["B"]:
        assert row["context"].startswith("[K1]\nQ: dates") and k["has_context"] is True
        assert "[SO-" not in row["candidate_answer"] and "retrieved sources" not in row["candidate_answer"]
        assert "[catatan: informasi tidak tersedia] tz." in row["candidate_answer"]
    assert set(export["key_items"][0]["text_sha256"]) == {"question", "accepted_answer", "candidate_answer", "context"}


def test_annotator_id_is_prefilled_per_sheet(sheet):
    from openpyxl import load_workbook

    _, pdir = sheet
    for name in ("annotator1", "annotator2"):
        ws = load_workbook(pdir / f"{name}.xlsx")["anotasi"]
        col = at.COLUMNS.index("annotator_id") + 1
        assert {ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)} == {name}


def _set(path, item_id, **values):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["anotasi"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == item_id:
            for name, v in values.items():
                row[at.COLUMNS.index(name)].value = v
    wb.save(path)


def _row_ids(export, condition):
    return [k["item_id"] for k in export["key_items"] if k["condition"] == condition]


def _problems_for(path, item_id):
    rows, problems = at.read_annotation_xlsx(path)
    return {r["item_id"]: r for r in rows}[item_id], [p for p in problems if f"({item_id})" in p]


VALID_B = dict(abstain="N", hallucinated="Y", severity="PENUH",
               wrong_claims='fabricated | parse_iso() exists | "[K1] A: use fromisoformat"',
               faithfulness_level="partly_supported", approach_match="different_valid",
               domain_familiarity="familiar", uncertain="N")


def test_valid_rows_and_parsed_wrong_claims(sheet):
    export, pdir = sheet
    b, a = _row_ids(export, "B")[0], _row_ids(export, "A")[0]
    _set(pdir / "annotator1.xlsx", b, **{**VALID_B, "wrong_claims": VALID_B["wrong_claims"] +
                                         "\ncontradicted | X is default | https://docs.python.org/3/library/x.html"})
    _set(pdir / "annotator1.xlsx", a, abstain="N", hallucinated="N", approach_match="same",
         domain_familiarity="not_familiar", uncertain="Y")
    row, problems = _problems_for(pdir / "annotator1.xlsx", b)
    assert problems == [] and row["category"] == "HALLUCINATED" and row["annotator_id"] == "annotator1"
    assert [w["type"] for w in row["wrong_claims"]] == ["fabricated", "contradicted"]
    assert row["faithfulness_level"] == "partly_supported" and row["has_context"] is True
    row, problems = _problems_for(pdir / "annotator1.xlsx", a)
    assert problems == [] and row["category"] == "NOT_HALLUCINATED" and row["uncertain"] == "Y"


@pytest.mark.parametrize("override, needle", [
    ({"wrong_claims": None}, "needs at least one wrong claim with evidence"),
    ({"wrong_claims": "fabricated | parse_iso() exists | I think so"}, "evidence must be a URL or a quoted passage"),
    ({"wrong_claims": "wrong | parse_iso() exists | https://x.org"}, "must be contradicted or fabricated"),
    ({"wrong_claims": "fabricated parse_iso https://x.org"}, "expected 'type | claim | evidence'"),
    ({"faithfulness_level": None}, "faithfulness_level is None"),
    ({"approach_match": "similar"}, "approach_match is 'similar'"),
    ({"domain_familiarity": None}, "domain_familiarity is None"),
    ({"uncertain": "maybe"}, "uncertain is 'MAYBE'"),
    ({"annotator_id": None}, "annotator_id is empty"),
])
def test_invalid_rows_are_rejected(sheet, override, needle):
    export, pdir = sheet
    b = _row_ids(export, "B")[0]
    _set(pdir / "annotator1.xlsx", b, **{**VALID_B, **override})
    row, problems = _problems_for(pdir / "annotator1.xlsx", b)
    assert row["category"] is None and any(needle in p for p in problems), problems


def test_condition_a_rows_must_leave_faithfulness_empty_and_not_hallucinated_rows_no_wrong_claims(sheet):
    export, pdir = sheet
    a = _row_ids(export, "A")[0]
    _set(pdir / "annotator1.xlsx", a, abstain="N", hallucinated="N", approach_match="same",
         faithfulness_level="all_supported", wrong_claims="fabricated | x | https://x.org",
         domain_familiarity="familiar", uncertain="N")
    _, problems = _problems_for(pdir / "annotator1.xlsx", a)
    text = "\n".join(problems)
    assert "faithfulness_level must be empty when the context cell is empty" in text
    assert "wrong_claims must be empty when hallucinated = N" in text


def test_old_format_sheet_is_refused(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "anotasi"
    ws.append(["item_id", "question", "accepted_answer", "candidate_answer", "abstain", "hallucinated", "severity",
               "notes"])
    wb.save(tmp_path / "old.xlsx")
    rows, problems = at.read_annotation_xlsx(tmp_path / "old.xlsx")
    assert rows == [] and "older export" in problems[0]
    assert json.dumps(at.COLUMNS)  # sanity
