"""Amendment 11: 20% intra-annotator re-annotation (>= 14 days after completion) and single-annotator
reference."""

import json
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest

import annotation_tool as at
from test_annotation_sheet_v3 import _items

T0 = datetime(2026, 11, 1, tzinfo=timezone.utc)


def _fill_all(path, flip=()):
    """Every row valid; items in `flip` get the opposite hallucinated label."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["anotasi"]
    c = {n: i for i, n in enumerate(at.COLUMNS)}
    for row in ws.iter_rows(min_row=2):
        qnum = int(str(row[c["question"]].value).split()[1])          # "Title 10xx ..."
        hall = "Y" if (qnum % 3 == 0) != (row[0].value in flip) else "N"
        row[c["abstain"]].value, row[c["hallucinated"]].value = "N", hall
        if hall == "Y":
            row[c["severity"]].value = "PENUH"
            row[c["wrong_claims"]].value = "fabricated | foo() | https://docs.example.org"
        row[c["approach_match"]].value = "same"
        if row[c["context"]].value:
            row[c["faithfulness_level"]].value = "all_supported"
        row[c["domain_familiarity"]].value, row[c["uncertain"]].value = "familiar", "N"
    wb.save(path)


@pytest.fixture
def env(tmp_path, monkeypatch):
    root, key_dir = tmp_path / "annotation", tmp_path / "keys"
    monkeypatch.setattr(at, "ANNOTATION_ROOT", root)
    export = at.build_export(_items(), {q: f"<p>Body {q}</p>" for q in range(1000, 1050)}, "calibration", 42)
    pdir = at.purpose_dir("calibration")
    at.write_annotation_xlsx(pdir / "annotator1.xlsx", export["annotator1"], "annotator1")
    at.write_annotation_xlsx(pdir / "annotator2.xlsx", export["annotator2"], "annotator2")
    key_dir.mkdir()
    key_path = key_dir / "calibration_key_seed42.json"
    key_path.write_text(json.dumps({"purpose": "calibration", "seed": 42, "items": export["key_items"]}))
    (pdir / "export_manifest.json").write_text(json.dumps({"key_file_name": key_path.name, "n_items_annotator2": 30,
                                                           "key_sha256": at.file_sha256(key_path)}))
    _fill_all(pdir / "annotator1.xlsx")
    return SimpleNamespace(pdir=pdir, key_dir=key_dir, export=export)


def _args(env, **kw):
    return SimpleNamespace(purpose="calibration", annotator="annotator1", key_dir=str(env.key_dir), seed=42,
                           override_min_days=None, force=False, **kw)


def test_selection_is_20_percent_deterministic_with_new_ids_and_order():
    ids = [f"CAL-{i:03d}" for i in range(1, 61)]
    pairs = at.select_reannotation(ids, "calibration", 42)
    assert len(pairs) == 12 and pairs == at.select_reannotation(list(reversed(ids)), "calibration", 42)
    assert [n for n, _ in pairs] == [f"CALR-{i:02d}" for i in range(1, 13)]
    assert [o for _, o in pairs] != sorted(o for _, o in pairs) and len({o for _, o in pairs}) == 12


def test_export_needs_a_completed_sheet_and_14_days(env):
    with pytest.raises(at.AnnotationError, match="never validated as complete"):
        at.cmd_export_reannotation(_args(env), now=T0)
    at.record_completion("calibration", "annotator1", env.pdir / "annotator1.xlsx", now=T0)
    with pytest.raises(at.AnnotationError, match=r"only 13\.0 days"):
        at.cmd_export_reannotation(_args(env), now=T0 + timedelta(days=13))
    assert at.cmd_export_reannotation(_args(env), now=T0 + timedelta(days=15)) == 0
    man = json.loads((env.pdir / "reannotation_annotator1_manifest.json").read_text())
    assert man["n_items"] == 12 and man["override_min_days"] is False and man["days_since_completion"] == 15.0
    assert (env.key_dir / man["map_file_name"]).exists() and not (env.pdir / man["map_file_name"]).exists()

    from openpyxl import load_workbook
    ws = load_workbook(env.pdir / "reannotation_annotator1.xlsx")["anotasi"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == [f"CALR-{i:02d}" for i in range(1, 13)]
    labels = [at.COLUMNS.index(c) for c in at.FILL_COLUMNS if c != "annotator_id"]
    assert all(r[i] is None for r in rows for i in labels)                       # no first-pass labels
    assert {r[at.COLUMNS.index("annotator_id")] for r in rows} == {"annotator1"}


def test_override_is_recorded_and_edited_sheet_is_refused(env):
    at.record_completion("calibration", "annotator1", env.pdir / "annotator1.xlsx", now=T0)
    override = SimpleNamespace(**{**vars(_args(env)), "override_min_days": "trial run"})
    at.cmd_export_reannotation(override, now=T0 + timedelta(days=2))
    man = json.loads((env.pdir / "reannotation_annotator1_manifest.json").read_text())
    assert man["override_min_days"] is True and man["override_reason"] == "trial run"
    _fill_all(env.pdir / "annotator1.xlsx", flip={"CAL-001"})                   # edited after completion
    with pytest.raises(at.AnnotationError, match="changed after it was recorded complete"):
        at.cmd_export_reannotation(SimpleNamespace(**{**vars(_args(env)), "force": True}), now=T0 + timedelta(days=20))


def test_intra_kappa_counts_the_changed_labels(env):
    at.record_completion("calibration", "annotator1", env.pdir / "annotator1.xlsx", now=T0)
    at.cmd_export_reannotation(_args(env), now=T0 + timedelta(days=14))
    pairs = json.loads(next(env.key_dir.glob("*reannotation_map*")).read_text())["pairs"]
    _fill_all(env.pdir / "reannotation_annotator1.xlsx")                      # same rule -> same labels ...
    first = {r["item_id"]: r for r in at.read_annotation_xlsx(env.pdir / "annotator1.xlsx")[0]}
    second = {r["item_id"]: r for r in at.read_annotation_xlsx(env.pdir / "reannotation_annotator1.xlsx")[0]}
    rep = at.intra_annotator(first, second, pairs)
    assert rep["n"] == 12 and rep["changes"] == [] and rep["observed_agreement_3cat"] == 1.0
    flip = pairs[0]["new_id"]
    _fill_all(env.pdir / "reannotation_annotator1.xlsx", flip={flip})          # ... except one
    assert at.cmd_intra_kappa(_args(env)) == 0
    out = json.loads((env.pdir / "intra_annotator_annotator1.json").read_text())
    assert [c["new_id"] for c in out["changes"]] == [flip] and out["n"] == 12


def test_single_annotator_reference_when_annotator2_sheet_is_empty(env):
    assert at.sheet_started(env.pdir / "annotator1.xlsx") and not at.sheet_started(env.pdir / "annotator2.xlsx")
    rc = at.cmd_finalize_reference(SimpleNamespace(purpose="calibration", key_dir=str(env.key_dir)))
    ref = json.loads((env.pdir / "reference_adjudicated.json").read_text())
    assert rc == 0 and ref["n_items"] == 60 and {it["source"] for it in ref["items"]} == {"annotator1"}
    assert ref["human_human"] is None
