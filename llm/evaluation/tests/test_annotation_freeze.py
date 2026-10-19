"""Amendment 12: freezing the validated annotation sheet, and the frozen-hash guard. Offline, temporary files."""

import json
from datetime import datetime, timezone

import pytest

import annotation_tool as at
import judge_panel_agreement as ag
from test_reannotation import T0, _args, _fill_all, env  # noqa: F401 -- fixture


def test_freeze_writes_record_copy_and_completion(env):
    rec = at.freeze_annotation("calibration", "annotator1", env.key_dir, now=T0)
    sheet = env.pdir / "annotator1.xlsx"
    assert rec["sha256"] == at.file_sha256(sheet) and rec["item_count"] == 60
    assert rec["note"] == "frozen before any judge run on calibration items" and rec["file_name"] == "annotator1.xlsx"
    copy = env.key_dir / "annotator1_frozen_20261101.xlsx"
    assert copy.exists() and at.file_sha256(copy) == rec["sha256"]
    assert json.loads((env.pdir / "annotator1_freeze.json").read_text()) == rec
    completion = json.loads((env.pdir / "annotator1_completed.json").read_text())
    assert completion["reannotation"] == "not_scheduled" and completion["sheet_sha256"] == rec["sha256"]
    assert at.assert_annotation_frozen("calibration", "annotator1")["sha256"] == rec["sha256"]
    # idempotent for the same sheet
    assert at.freeze_annotation("calibration", "annotator1", env.key_dir, now=T0) == rec


def test_invalid_sheet_is_not_frozen(env):
    from openpyxl import load_workbook
    path = env.pdir / "annotator1.xlsx"
    wb = load_workbook(path)
    wb["anotasi"].cell(row=2, column=at.COLUMNS.index("uncertain") + 1).value = None
    wb.save(path)
    with pytest.raises(at.AnnotationError, match="not valid"):
        at.freeze_annotation("calibration", "annotator1", env.key_dir, now=T0)
    assert not (env.pdir / "annotator1_freeze.json").exists()


def test_changed_sheet_fails_the_guard_and_cannot_be_refrozen(env):
    at.freeze_annotation("calibration", "annotator1", env.key_dir, now=T0)
    _fill_all(env.pdir / "annotator1.xlsx", flip={env.export["annotator1"][0]["item_id"]})
    with pytest.raises(at.AnnotationError, match="changed after it was frozen"):
        at.assert_annotation_frozen("calibration", "annotator1")
    with pytest.raises(at.AnnotationError, match="different version"):
        at.freeze_annotation("calibration", "annotator1", env.key_dir, now=T0)


def test_guard_without_freeze(env):
    with pytest.raises(at.AnnotationError, match="has not been frozen"):
        at.assert_annotation_frozen("calibration", "annotator1")


def test_agreement_script_refuses_unfrozen_or_changed_sheet(env, tmp_path):
    args = ["--purpose", "calibration", "--judge-output", str(tmp_path / "j.jsonl"),
            "--panel-output", str(tmp_path / "p.jsonl"), "--key-dir", str(env.key_dir)]
    assert ag.main(args) == 1                      # not frozen
    at.freeze_annotation("calibration", "annotator1", env.key_dir, now=T0)
    _fill_all(env.pdir / "annotator1.xlsx", flip={env.export["annotator1"][0]["item_id"]})
    assert ag.main(args) == 1                      # changed after freezing


def test_reannotation_refused_when_not_scheduled(env):
    at.freeze_annotation("calibration", "annotator1", env.key_dir, now=T0)
    with pytest.raises(at.AnnotationError, match="not scheduled"):
        at.cmd_export_reannotation(_args(env), now=datetime(2027, 1, 1, tzinfo=timezone.utc))


def test_export_v2_sheet_copies_frozen_v1_and_adds_readonly_columns(env):
    from openpyxl import load_workbook
    at.freeze_annotation("calibration", "annotator1", env.key_dir, now=T0)
    v1 = env.pdir / "annotator1.xlsx"
    v1_sha = at.file_sha256(v1)
    ids = [r["item_id"] for r in env.export["annotator1"]]
    dates = {i: ("2019-03-01", "2019-03-02") for i in ids}
    dates[ids[0]] = (None, None)
    m = at.export_v2_sheet("calibration", "annotator1", dates, ids[1:3], now=T0)
    assert at.file_sha256(v1) == v1_sha                                   # v1 untouched
    out = env.pdir / "annotator1_v2.xlsx"
    ws = load_workbook(out)["anotasi"]
    header = [c.value for c in ws[1]]
    assert header == at.COLUMNS + at.V2_COLUMNS
    rows = {r[0]: dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    assert all(rows[i]["v2_change_reason"] is None for i in ids)          # nothing pre-filled
    assert [i for i in ids if rows[i]["time_rule_check"] == "Y"] == ids[1:3]
    assert rows[ids[0]]["question_date"] is None and rows[ids[1]]["question_date"] == "2019-03-01"
    # the v1 part is identical and still valid
    a, pa = at.read_annotation_xlsx(v1)
    b, pb = at.read_annotation_xlsx(out)
    assert pa == pb == [] and [{k: v for k, v in r.items() if k != "row"} for r in a] == \
        [{k: v for k, v in r.items() if k != "row"} for r in b]
    assert m["source_sha256"] == v1_sha and m["n_question_date_missing"] == 1
    assert m["expected_readonly"][ids[1]] == {"question_date": "2019-03-01", "reference_date": "2019-03-02",
                                              "time_rule_check": "Y"}
    with pytest.raises(at.AnnotationError, match="exported once"):
        at.export_v2_sheet("calibration", "annotator1", dates, ids[1:3], now=T0)


def test_export_v2_needs_the_frozen_sheet(env):
    with pytest.raises(at.AnnotationError, match="not been frozen"):
        at.export_v2_sheet("calibration", "annotator1", {}, [], now=T0)
